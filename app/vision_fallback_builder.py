import json
import logging
import re
from typing import Any, Dict, Optional
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from app.color_utils import (
    extract_global_design_tokens,
    hex_to_pptx_color as _hex_to_pptx_color,
    hex_to_rgb_tuple as _hex_to_rgb_tuple,
    parse_color_value,
)
from app.genai_client import DEFAULT_GEMINI_MODEL, get_genai_client
from app.pptx_native_builders import (
    apply_semantic_styles_to_table,
    audit_and_resolve_slide_collisions,
    build_slide_from_geometry,
    build_styled_native_chart,
    build_styled_native_table,
    embed_extracted_slide_images,
    ensure_slide_canvas_background,
    ensure_slide_typography_consistency,
    refine_card_accent_bars,
)
from app.schemas import SlideNativeDecomposition, SlideNativeElement, TextRun

logger = logging.getLogger(__name__)


def decompose_slide_image_with_vision(
    image_path: str,
    model_name: str = DEFAULT_GEMINI_MODEL,
) -> Optional[SlideNativeDecomposition]:
  """Decomposes a rendered 16:9 slide screenshot into 100% native slide elements using Gemini Vision."""
  try:
    from google import genai
    from google.genai import types

    with open(image_path, "rb") as f:
      image_bytes = f.read()

    client = get_genai_client()

    prompt = """
    You are an expert Vision-to-Native-Slide Reverse Engineer.
    Analyze this 16:9 widescreen slide image and decompose it into 100% editable native PowerPoint slide elements.

    [STRICT MANDATES]
    1. ZERO-OMISSION OCR: Transcribe 100% of all text, titles, subtitles, bullets, table cells, metric numbers, badges, and footnotes verbatim without dropping or summarizing anything.
    2. NORMALIZED COORDINATES: Provide exact bounding box percentages (left_pct, top_pct, width_pct, height_pct from 0.0 to 100.0) for each element so they align precisely with the screenshot visual layout.
    3. ELEMENT TYPES:
       - 'CONTAINER': Background cards, rounded surface tiles, banners. (Specify bg_color_hex, border_color_hex, border_radius_pt).
       - 'TEXT_BOX': Independent titles, paragraphs, bullet lists, metric callouts. (Specify font_size_pt, is_bold, font_color_hex, alignment).
       - 'TABLE': Tabular matrices (provide table_matrix as 2D list of row strings, table_header_bg_hex).
       - 'DIVIDER': Thin horizontal/vertical separating lines.
       - 'BADGE': Small category tags/pills.
    4. DESIGN TOKENS: Extract exact HEX colors for background canvas, text colors, container surfaces, and border strokes.
    """

    response = client.models.generate_content(
        model=model_name,
        contents=[
            types.Part.from_bytes(data=image_bytes, mime_type="image/png"),
            prompt,
        ],
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=SlideNativeDecomposition,
            temperature=0.1,
        ),
    )

    if response.text:
      return SlideNativeDecomposition.model_validate_json(response.text)
  except Exception as e:
    logger.warning("[Vision Decomposition Error] %s: %s", type(e).__name__, e)
    return None


def generate_native_slide_builder_code(
    image_path: str,
    slide_dom_data: Optional[Dict[str, Any]] = None,
    slide_geometry: Optional[Dict[str, Any]] = None,
    model_name: str = DEFAULT_GEMINI_MODEL,
    max_retries: int = 2,
    design_tokens: Optional[Dict[str, Any]] = None,
    slide_index: Optional[int] = None,
    total_slides: Optional[int] = None,
) -> Optional[str]:
  """Uses Gemini Vision fusing 16:9 screenshot and semantic HTML DOM to generate 100% native slide builder code."""
  from app.slide_fidelity_critic import audit_slide_fidelity, run_code_on_test_slide
  try:
    with open(image_path, "rb") as f:
      image_bytes = f.read()

    client = get_genai_client()

    tokens = design_tokens or extract_global_design_tokens(first_slide_geometry=slide_geometry)
    b_r, b_g, b_b = _hex_to_rgb_tuple(tokens.get("brand_color_hex", "#2563EB"), (37, 99, 235))
    t_r, t_g, t_b = _hex_to_rgb_tuple(tokens.get("text_main_hex", "#111115"), (17, 17, 21))
    m_r, m_g, m_b = _hex_to_rgb_tuple(tokens.get("text_muted_hex", "#64748B"), (100, 116, 139))
    bg_r, bg_g, bg_b = _hex_to_rgb_tuple(tokens.get("bg_color_hex", "#FFFFFF"), (255, 255, 255))
    if slide_geometry and slide_geometry.get("slideBgColor"):
      s_bg_col = parse_color_value(slide_geometry["slideBgColor"], default=RGBColor(bg_r, bg_g, bg_b))
      bg_r, bg_g, bg_b = s_bg_col[0], s_bg_col[1], s_bg_col[2]
    is_dark = (bg_r * 0.299 + bg_g * 0.587 + bg_b * 0.114) < 128
    font_family = tokens.get("font_name", "Pretendard")
    slide_label = f"Slide {slide_index} of {total_slides}" if slide_index and total_slides else "Slide"

    # Measured DOM layout coordinates snippet (strip large base64 dataUrl from charts)
    measured_layout_snippet = ""
    if slide_geometry:
      geom_clean = {}
      for k, v in slide_geometry.items():
        if not v:
          continue
        if k == "charts" and isinstance(v, list):
          geom_clean[k] = [{ck: cv for ck, cv in ch.items() if ck != "dataUrl"} for ch in v]
        else:
          geom_clean[k] = v
      geom_json = json.dumps(geom_clean, ensure_ascii=False, indent=2)
      measured_layout_snippet = f"""
[EXACT MEASURED BROWSER LAYOUT GEOMETRY (PIXEL-PERFECT INCHES)]
Use these exact coordinates and styling values measured directly from Chrome:
{geom_json}
"""

    # Semantic context if parsed DOM is available
    semantic_context = ""
    has_tables = False
    table_snippet = ""
    if slide_dom_data:
      has_tables = len(slide_dom_data.get("tables", [])) > 0
      semantic_json = json.dumps(
          {k: v for k, v in slide_dom_data.items() if k != "raw_html"},
          ensure_ascii=False,
          indent=2,
      )
      semantic_context = f"""
[PARSED SEMANTIC HTML TRUTH]
{semantic_json}

[SLIDE ORIGINAL HTML SOURCE]
```html
{slide_dom_data.get('raw_html', '')}
```
"""
      if has_tables:
        t_data_json = json.dumps(slide_dom_data["tables"][0], ensure_ascii=False)
        is_slide_dark = is_dark or (tokens.get("text_main_hex") == "#F8FAFC")
        if is_slide_dark:
          table_style_rules = """- Dark theme header background `RGBColor(15, 23, 42)` with muted light text `RGBColor(148, 163, 184)`.
- Alternating dark row fills (`RGBColor(30, 41, 59)` and `RGBColor(22, 30, 46)`).
- Subtotal accent blue row fill (`RGBColor(28, 48, 79)`) with light blue text `RGBColor(147, 197, 253)`.
- Total row deep navy fill (`RGBColor(15, 23, 42)`) with bright white bold text `RGBColor(255, 255, 255)`.
- High contrast readable cell text: default body text `RGBColor(248, 250, 252)`. Never use dark text on dark cells!"""
        else:
          table_style_rules = f"""- Light header background `RGBColor(241, 245, 249)` with bold accent text `RGBColor({b_r}, {b_g}, {b_b})`.
- Alternating row fills (`#FFFFFF` and `#F8FAFC`).
- Semantic cell alert colors:
  * Red alert text (`#D93025`): `RGBColor(217, 48, 37)`
  * Gold alert text (`#B48000`): `RGBColor(180, 128, 0)`
  * Green alert text (`#008744`): `RGBColor(0, 135, 68)`
  * Parameter keys in Column 1: `bold = True`"""
        table_snippet = f"""
[TABLE BUILDER INSTRUCTIONS - 100% PRECISE DATA & STYLING]
This slide contains a data table. You can use the built-in helper function available in scope:
`build_styled_native_table(slide, table_data, left, top, width, height, col_ratios=None)`
Example:
```python
t_data = {t_data_json}
build_styled_native_table(slide, t_data, Inches(0.8), Inches(1.95), Inches(11.733), Inches(4.5))
```
Or you may build the table natively using `slide.shapes.add_table`, adhering strictly to:
{table_style_rules}
"""

    # Build structured Slide Manifest for the slide-building AI
    manifest_lines = []
    manifest_lines.append(f"Canvas: 16:9 Widescreen (13.333\" x 7.5\"), Background = RGBColor({bg_r}, {bg_g}, {bg_b}) (HEX: #{bg_r:02X}{bg_g:02X}{bg_b:02X}, is_dark={is_dark})")
    manifest_lines.append(f"Global Font Family: {font_family}")
    
    if slide_geometry:
      header_badges = slide_geometry.get("headerBadges", [])
      if header_badges:
        for hb_idx, hb in enumerate(header_badges):
          r = hb.get("rect", {})
          st = hb.get("styles", {})
          manifest_lines.append(f"Header Category Pill/Badge #{hb_idx+1}: \"{hb.get('text')}\" at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')}), fontSize={st.get('fontSizePt', 10.0)}pt, color={st.get('color')}, bg={st.get('backgroundColor')}")
      elif slide_geometry.get("tag"):
        t = slide_geometry["tag"]
        r = t.get("rect", {})
        st = t.get("styles", {})
        manifest_lines.append(f"Header Category Pill/Badge: \"{t.get('text')}\" at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')}), fontSize={st.get('fontSizePt', 10.0)}pt, color={st.get('color')}, bg={st.get('backgroundColor')}")
      
      if slide_geometry.get("title"):
        t = slide_geometry["title"]
        r = t.get("rect", {})
        st = t.get("styles", {})
        manifest_lines.append(f"Main Slide Title: \"{t.get('text')}\" at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')}), fontSize={st.get('fontSizePt', 24.0)}pt, bold=True, color={st.get('color')}")
      
      if slide_geometry.get("sub"):
        s = slide_geometry["sub"]
        r = s.get("rect", {})
        st = s.get("styles", {})
        manifest_lines.append(f"Subtitle / Description: \"{s.get('text')}\" at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')}), fontSize={st.get('fontSizePt', 13.0)}pt, color={st.get('color')}")
      
      cards = slide_geometry.get("cards", [])
      if cards:
        manifest_lines.append(f"\nVisual Layout Containers ({len(cards)} containers detected):")
        for c_idx, c in enumerate(cards):
          r = c.get("rect", {})
          st = c.get("styles", {})
          is_rnd = st.get("isRounded", st.get("borderRadius", 0) >= 6)
          c_title = c.get("title") or "Container"
          manifest_lines.append(f"  Container #{c_idx+1}: [{c_title}] at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')})")
          manifest_lines.append(f"    - Shape: {'MSO_SHAPE.ROUNDED_RECTANGLE' if is_rnd else 'MSO_SHAPE.RECTANGLE'} (borderRadius={st.get('borderRadius', 0)}px)")
          manifest_lines.append(f"    - Fill: {st.get('backgroundColor')}, Border: {st.get('borderColor')}")
          if c.get("hasTopAccent"):
            manifest_lines.append(f"    - Top Accent Bar: color={c.get('topAccentColor')}")
          if c.get("hasLeftAccent"):
            manifest_lines.append(f"    - Left Accent Bar: color={c.get('leftAccentColor')}")
          if c.get("isBridge") and c.get("bridge"):
            manifest_lines.append(f"    - Bridge Connector Banner: \"{c['bridge'].get('text')}\" (arrow={c['bridge'].get('arrow')})")
          if c.get("isAttachCard") and c.get("attachData"):
            ad = c["attachData"]
            manifest_lines.append(f"    - Attach Card Details: target=\"{ad.get('target')}\", mode=\"{ad.get('mode')}\", mech=\"{ad.get('mech')}\", desc=\"{ad.get('desc')}\"")
          if c.get("pipelineSteps"):
            manifest_lines.append(f"    - Pipeline Steps ({len(c['pipelineSteps'])} steps):")
            for ps in c["pipelineSteps"]:
              manifest_lines.append(f"      * Step {ps.get('stepNum')}: \"{ps.get('title')}\" - {ps.get('desc')}")
          if c.get("subCards"):
            manifest_lines.append(f"    - Nested Sub-Cards ({len(c['subCards'])} cards):")
            for sc in c["subCards"]:
              manifest_lines.append(f"      * [{sc.get('head', '')}] badge='{sc.get('badge', '')}', items={sc.get('items', [])}")
          if c.get("flows"):
            manifest_lines.append(f"    - Inline Flows:")
            for fl in c["flows"]:
              node_texts = [n.get('text', '') if isinstance(n, dict) else str(n) for n in fl.get('nodes', [])]
              manifest_lines.append(f"      * Sequence: {' -> '.join(node_texts)} (arrow: {fl.get('arrow')})")
          if c.get("fItems"):
            manifest_lines.append(f"    - Feature Items with Icon Badges:")
            for fi in c["fItems"]:
              manifest_lines.append(f"      * [{fi.get('icon')}] \"{fi.get('text')}\"")
          for b in c.get("badges", []):
            manifest_lines.append(f"    - Inner Badge: \"{b.get('text')}\"")
          for p in c.get("paragraphs", []):
            manifest_lines.append(f"    - Text Item: \"{p.get('text')}\" (color={p.get('styles', {}).get('color')})")
          for cb in c.get("codeBlocks", []):
            manifest_lines.append(f"    - Code / Calculation Block: \"{cb.get('text')}\"")
      
      charts = slide_geometry.get("charts", [])
      if charts:
        manifest_lines.append(f"\nData Charts ({len(charts)} detected):")
        for ch_idx, ch in enumerate(charts):
          r = ch.get("rect", {})
          cfg = ch.get("chartConfig") or {}
          manifest_lines.append(f"  Chart #{ch_idx+1}: type={cfg.get('type', 'unknown')} at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')})")

      tables = slide_geometry.get("tables", [])
      if tables:
        manifest_lines.append(f"\nData Tables ({len(tables)} detected):")
        for t_idx, tbl in enumerate(tables):
          r = tbl.get("rect", {})
          manifest_lines.append(f"  Table #{t_idx+1}: {tbl.get('num_rows')} rows x {tbl.get('num_cols')} columns at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')})")
      
      hl = slide_geometry.get("highlightBoxes", [])
      if hl:
        manifest_lines.append(f"\nCallout / Highlight Boxes ({len(hl)} detected):")
        for h_idx, h in enumerate(hl):
          r = h.get("rect", {})
          manifest_lines.append(f"  Callout #{h_idx+1}: \"{h.get('text')}\" at (left={r.get('left')}, top={r.get('top')}, width={r.get('width')}, height={r.get('height')})")
          manifest_lines.append(f"    - Fill: {h.get('styles', {}).get('backgroundColor')}, Border: {h.get('styles', {}).get('borderColor')}")
      
      fn = slide_geometry.get("footnotes", [])
      if fn:
        manifest_lines.append(f"\nFootnotes ({len(fn)} detected):")
        for f_idx, f in enumerate(fn):
          r = f.get("rect", {})
          manifest_lines.append(f"  Footnote #{f_idx+1}: \"{f.get('text')}\" at (left={r.get('left')}, top={r.get('top')})")

    manifest_context = "\n".join(manifest_lines)

    base_prompt = f"""
You are an expert PowerPoint Engineer specializing in high-fidelity native PPTX recreation.
Analyze this 16:9 widescreen slide image ({slide_label}) and the semantic HTML structure & style manifest below.
Generate a Python function `def build_slide(prs, slide):` using `python-pptx` that reconstructs this slide into 100% native, fully editable shapes, text boxes, charts, and tables.

[AVAILABLE IN EXECUTION SCOPE]
- `prs`, `slide`, `Inches`, `Pt`, `RGBColor`, `MSO_SHAPE`, `PP_ALIGN`, `MSO_ANCHOR`
- `build_styled_native_table(slide, table_data, left, top, width, height, col_ratios=None)`
- `build_styled_native_chart(slide, chart_data, font_name="{font_family}", is_dark={is_dark}, slide_bg_rgb=({bg_r}, {bg_g}, {bg_b}))`

[AUTHENTIC SLIDE STRUCTURE & STYLE MANIFEST (MEASURED GROUND TRUTH)]
{manifest_context}

[STRICT INSTRUCTIONS FOR THE SLIDE BUILDER AI]
1. ZERO HARDCODING (DYNAMIC RECONSTRUCTION):
   - DO NOT assume any fixed master grid or fixed template coordinates.
   - Faithfully reconstruct the exact layout, containers, charts, tables, and elements defined in the MANIFEST above.
   - If there are 4 metric cards, build 4 metric cards. If there are 2 comparison containers, build 2 comparison containers. If there is a table and 3 bottom cards, build them exactly in that order!
2. EXACT SHAPES & CORNER RADII (MSO_SHAPE.ROUNDED_RECTANGLE):
   - Any container or box with `isRounded=True` or `borderRadius >= 6px` MUST be created using `MSO_SHAPE.ROUNDED_RECTANGLE`. Never create sharp rectangular boxes where rounded cards exist!
   - For cards with top accent colors, add an inner rounded capsule pill (`MSO_SHAPE.ROUNDED_RECTANGLE` with `adjustments[0] = 0.5`, height ~0.06 in, inset inside the card: `left = card_left + Inches(0.20)`, `top = card_top + Inches(0.08)`, `width = card_width - Inches(0.40)`).
3. 100% FAITHFUL STYLING & HIGH CONTRAST:
   - Canvas background MUST be full-bleed 16:9 rectangle: `slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))` with fill `RGBColor({bg_r}, {bg_g}, {bg_b})` and no border line.
   - For every container and text run, apply the exact RGB colors from the manifest.
   - Dark theme rule: On dark containers (e.g. #0F172A, #1E293B), all text MUST be high contrast (white or light tint like RGBColor(248, 250, 252) or light slate RGBColor(148, 163, 184)). NEVER use dark text on dark cards!
4. ZERO OMISSION MANDATE (100% Verbatim):
   - Transcribe 100% of all Korean and English titles, metrics, card items, bullet points, nested spec/calculation blocks, callout boxes (💡핵심 요약, 📌참고, 💡결론), and bottom footnotes (* ...) verbatim!
   - Every single bullet point must be preserved without skipping or summarizing.
5. ZERO CLIPPING & DYNAMIC HEIGHT BUDGETING:
   - Set `tf.word_wrap = True` and give text boxes adequate height so text never overflows or gets clipped.
   - When a table is present, place bottom cards below the table without overlapping.
{table_snippet}
{measured_layout_snippet}
{semantic_context}

Output ONLY executable Python code defining:
def build_slide(prs, slide):
    # implementation
"""

    current_prompt = base_prompt
    best_code = None

    for attempt in range(1, max_retries + 1):
      code = ""
      for api_try in range(3):
        try:
          response = client.models.generate_content(
              model=model_name,
              contents=[
                  types.Part.from_bytes(data=image_bytes, mime_type="image/png"),
                  current_prompt,
              ],
              config=types.GenerateContentConfig(
                  thinking_config=types.ThinkingConfig(thinking_level=types.ThinkingLevel.MEDIUM),
                  temperature=0.1,
              ),
          )
          code = response.text or ""
          if code:
            break
        except Exception as api_err:
          err_str = str(api_err)
          logger.warning("[Vision CodeGen API Try %d] Transient error: %s", api_try + 1, api_err)
          if (
              "reauthentication is needed" in err_str.lower()
              or "defaultcredentialserror" in type(api_err).__name__.lower()
          ):
            return None
          if api_try < 2:
            import time
            time.sleep(3 * (api_try + 1))
      if not code:
        continue
      best_code = code

      # 1. Structural runtime safety check
      success, test_slide, test_prs, exec_err = run_code_on_test_slide(code)
      if not success:
        logger.warning("[Slide CodeGen Attempt %d] Execution error: %s", attempt, exec_err)
        current_prompt = (
            base_prompt
            + f"\n\n[CRITICAL ERROR IN PREVIOUS ATTEMPT]:\nExecution failed with error: {exec_err}\nFix syntax, imports, and variables, and regenerate complete code."
        )
        continue

      # 2. Slide Fidelity & Zero-Omission Critic Audit
      report = audit_slide_fidelity(
          original_image_input=image_bytes,
          generated_code=code,
          slide=test_slide,
          model_name=model_name,
      )

      if report.passed and report.score >= 7:
        logger.info("[Slide Fidelity Audit PASSED on Attempt %d] Score: %d/10", attempt, report.score)
        return code
      else:
        logger.warning(
            "[Slide Fidelity Audit REJECTED on Attempt %d] Score: %d/10. Issues: %s, Missing: %s",
            attempt,
            report.score,
            report.design_and_layout_issues,
            report.missing_or_truncated_text,
        )
        current_prompt = (
            base_prompt
            + "\n\n[PREVIOUS ATTEMPT REJECTED BY FIDELITY AUDITOR - CORRECTIVE DIRECTIVES]:\n"
            + f"- Score: {report.score}/10 (Must be >= 7)\n"
            + (f"- Missing/Truncated text: {', '.join(report.missing_or_truncated_text)}\n" if report.missing_or_truncated_text else "")
            + (f"- Design/Layout defects: {', '.join(report.design_and_layout_issues)}\n" if report.design_and_layout_issues else "")
            + f"- Instructions: {report.corrective_instructions}\n"
            + "Please fix ALL issues and regenerate the complete, flawless Python code."
        )

    return best_code
  except Exception as e:
    logger.warning("[Vision CodeGen Error] %s: %s", type(e).__name__, e)
    return None


def build_native_slide_from_code(
    slide,
    prs: Presentation,
    code_str: Optional[str],
    fallback_img_path: str,
    design_tokens: Optional[Dict[str, Any]] = None,
    slide_dom_data: Optional[Dict[str, Any]] = None,
    slide_geometry: Optional[Dict[str, Any]] = None,
):
  """Executes generated python-pptx builder code to populate native slide, with geometry native builder and 2-tier fallback."""
  font_family = (design_tokens or {}).get("font_name", "Pretendard")
  text_color_hex = (design_tokens or {}).get("text_main_hex", "#111115")
  default_text_color = _hex_to_pptx_color(text_color_hex, default=(17, 17, 21))
  brand_color = _hex_to_pptx_color((design_tokens or {}).get("brand_color_hex"), default=(37, 99, 235))

  # Authoritative slide background & theme detection
  geom_bg = (slide_geometry or {}).get("slideBgColor")
  token_bg = (design_tokens or {}).get("bg_color_hex")
  slide_bg_rgb = parse_color_value(geom_bg or token_bg, default=RGBColor(255, 255, 255))
  is_dark = (slide_bg_rgb[0] * 0.299 + slide_bg_rgb[1] * 0.587 + slide_bg_rgb[2] * 0.114) < 128
  if default_text_color == RGBColor(17, 17, 21) and is_dark:
    default_text_color = RGBColor(241, 245, 249)

  # Pre-guarantee full-bleed 16:9 canvas background shape
  ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)

  if not code_str:
    if slide_geometry:
      build_slide_from_geometry(
          slide,
          slide_geometry,
          font_name=font_family,
          brand_color_rgb=brand_color,
          screenshot_path=fallback_img_path,
      )
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      refine_card_accent_bars(slide)
      audit_and_resolve_slide_collisions(slide)
      return

    decomp = decompose_slide_image_with_vision(fallback_img_path)
    if decomp and decomp.elements:
      build_native_pptx_slide(slide, decomp, prs, fallback_img_path)
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      return

    slide.shapes.add_picture(
        fallback_img_path,
        Inches(0),
        Inches(0),
        width=prs.slide_width,
        height=prs.slide_height,
    )
    return

  clean_code = code_str
  if "```python" in clean_code:
    clean_code = clean_code.split("```python", 1)[1].split("```", 1)[0].strip()
  elif "```" in clean_code:
    clean_code = clean_code.split("```", 1)[1].split("```", 1)[0].strip()

  scope = {
      "prs": prs,
      "slide": slide,
      "Presentation": Presentation,
      "Inches": Inches,
      "Pt": Pt,
      "RGBColor": RGBColor,
      "MSO_SHAPE": MSO_SHAPE,
      "PP_ALIGN": PP_ALIGN,
      "MSO_ANCHOR": MSO_ANCHOR,
      "build_styled_native_table": build_styled_native_table,
      "build_styled_native_chart": build_styled_native_chart,
  }

  try:
    exec(clean_code, scope)
    if "build_slide" in scope:
      func = scope["build_slide"]
      import inspect
      sig = inspect.signature(func)
      if len(sig.parameters) == 1:
        func(slide)
      else:
        func(prs, slide)

      # Ensure canvas background exists even if code wiped or drew over it
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)

      # Safety Net: Table construction and styling
      tbl_dom = (slide_geometry.get("tables", [{}])[0] if slide_geometry and slide_geometry.get("tables")
                 else (slide_dom_data.get("tables", [{}])[0] if slide_dom_data and slide_dom_data.get("tables") else None))
      if tbl_dom:
        has_table_shape = any(s.has_table for s in slide.shapes)
        if not has_table_shape and tbl_dom.get("rows"):
          logger.info("[Safety Net] Constructing table via deterministic builder because it was missing in executed code.")
          tbl_r = tbl_dom.get("rect", {})
          t_left = Inches(tbl_r.get("left", 0.6))
          t_top = Inches(tbl_r.get("top", 1.5))
          t_w = Inches(min(tbl_r.get("width", 12.0), 12.2))
          t_h = Inches(tbl_r.get("height", 3.0))
          tbl_copy = dict(tbl_dom)
          tbl_copy["is_dark"] = is_dark
          tbl_copy["slide_bg_rgb"] = (slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2])
          build_styled_native_table(slide, tbl_copy, t_left, t_top, t_w, t_h, font_name=font_family)
        else:
          for s in slide.shapes:
            if s.has_table:
              apply_semantic_styles_to_table(s, tbl_dom, font_name=font_family, is_dark=is_dark, slide_bg_rgb=(slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2]))

      # Safety Net: Chart construction
      if slide_geometry and slide_geometry.get("charts"):
        has_chart_shape = any(getattr(s, "has_chart", False) for s in slide.shapes)
        if not has_chart_shape:
          for ch in slide_geometry["charts"]:
            build_styled_native_chart(
                slide,
                ch,
                font_name=font_family,
                is_dark=is_dark,
                slide_bg_rgb=(slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2]),
            )

      # Safety Net: Zero Omission for SVG/IMG visuals, Footnotes, Callouts, and Bottom Cards
      if slide_geometry:
        embed_extracted_slide_images(slide, slide_geometry, fallback_img_path)

        slide_text_corpus = re.sub(r"\s+", " ", " ".join([
            p.text for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs
        ] + [
            c.text for s in slide.shapes if s.has_table for row in s.table.rows for c in row.cells
        ])).lower()

        # Check highlight boxes
        for hl in slide_geometry.get("highlightBoxes", []):
          hl_norm = re.sub(r"^[^a-zA-Z0-9가-힣]+", "", re.sub(r"\s+", " ", hl.get("text", "")).strip()).lower()
          hl_snippet = hl_norm[:18]
          hl_r = hl.get("rect", {})
          hl_l_in = float(hl_r.get("left", 0.8))
          hl_t_in = float(hl_r.get("top", 6.5))
          has_nearby_shape = any(
              s.has_text_frame
              and s.text_frame.text.strip()
              and abs((s.left.inches if hasattr(s.left, "inches") else s.left / 914400.0) - hl_l_in) < 0.35
              and abs((s.top.inches if hasattr(s.top, "inches") else s.top / 914400.0) - hl_t_in) < 0.35
              for s in slide.shapes
          )
          if hl_snippet and hl_snippet not in slide_text_corpus and not has_nearby_shape:
            logger.info("[Safety Net] Appending omitted highlight box: %s...", hl["text"][:30])
            hl_w = Inches(hl_r.get("width", 11.733))
            hl_l = Inches(hl_l_in)
            hl_t = Inches(hl_t_in)
            hl_h = Inches(max(0.40, hl_r.get("height", 0.42)))
            hl_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, hl_l, hl_t, hl_w, hl_h)
            hl_shape.fill.solid()
            hl_bg = parse_color_value(hl["styles"]["backgroundColor"], RGBColor(30, 41, 59) if is_dark else RGBColor(241, 245, 249), bg_blend_rgb=(slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2]))
            hl_border = parse_color_value(hl["styles"]["borderColor"], RGBColor(59, 130, 246) if is_dark else brand_color)
            hl_shape.fill.fore_color.rgb = hl_bg
            hl_shape.line.color.rgb = hl_border
            tf = hl_shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = hl["text"]
            p.font.name = font_family
            p.font.size = Pt(11.0)
            p.font.bold = True
            p.font.color.rgb = parse_color_value(hl["styles"]["color"], RGBColor(147, 197, 253) if is_dark else brand_color)

        # Check footnotes
        for fn in slide_geometry.get("footnotes", []):
          fn_norm = re.sub(r"^[^a-zA-Z0-9가-힣]+", "", re.sub(r"\s+", " ", fn.get("text", "")).strip()).lower()
          fn_snippet = fn_norm[:18]
          fn_r = fn.get("rect", {})
          fn_l_in = float(fn_r.get("left", 0.8))
          fn_t_in = float(fn_r.get("top", 7.0))
          has_nearby_fn = any(
              s.has_text_frame
              and s.text_frame.text.strip()
              and abs((s.left.inches if hasattr(s.left, "inches") else s.left / 914400.0) - fn_l_in) < 0.45
              and abs((s.top.inches if hasattr(s.top, "inches") else s.top / 914400.0) - fn_t_in) < 0.25
              for s in slide.shapes
          )
          if fn_snippet and fn_snippet not in slide_text_corpus and not has_nearby_fn:
            logger.info("[Safety Net] Appending omitted footnote: %s...", fn["text"][:30])
            fn_l = Inches(fn_l_in)
            fn_t = Inches(fn_t_in)
            fn_w = Inches(fn_r.get("width", 11.733))
            fn_h = Inches(max(0.25, fn_r.get("height", 0.25)))
            fn_box = slide.shapes.add_textbox(fn_l, fn_t, fn_w, fn_h)
            tf = fn_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = fn["text"]
            p.font.name = font_family
            p.font.size = Pt(9.5)
            p.font.color.rgb = parse_color_value(fn["styles"].get("color"), RGBColor(148, 163, 184) if is_dark else RGBColor(100, 116, 139))


      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      refine_card_accent_bars(slide)
      audit_and_resolve_slide_collisions(slide)
    elif len(slide.shapes) > 0:
      if slide_geometry:
        embed_extracted_slide_images(slide, slide_geometry, fallback_img_path)
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      refine_card_accent_bars(slide)
      audit_and_resolve_slide_collisions(slide)
    else:
      logger.warning("[Warn] 'build_slide' function not found and no shapes added.")
      if slide_geometry:
        build_slide_from_geometry(
            slide,
            slide_geometry,
            font_name=font_family,
            brand_color_rgb=brand_color,
            screenshot_path=fallback_img_path,
        )
        ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
        ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
        refine_card_accent_bars(slide)
        audit_and_resolve_slide_collisions(slide)
        return
      decomp = decompose_slide_image_with_vision(fallback_img_path)
      if decomp and decomp.elements:
        build_native_pptx_slide(slide, decomp, prs, fallback_img_path)
        ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
        ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
        audit_and_resolve_slide_collisions(slide)
      else:
        slide.shapes.add_picture(
            fallback_img_path,
            Inches(0),
            Inches(0),
            width=prs.slide_width,
            height=prs.slide_height,
        )
  except Exception as e:
    logger.warning("[Execution Error in generated slide code]: %s", e)
    if slide_geometry:
      logger.info("[Fallback] Building native slide from exact DOM geometry on execution error.")
      build_slide_from_geometry(
          slide,
          slide_geometry,
          font_name=font_family,
          brand_color_rgb=brand_color,
          screenshot_path=fallback_img_path,
      )
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      refine_card_accent_bars(slide)
      audit_and_resolve_slide_collisions(slide)
      return
    # If DOM data has table, construct native table slide
    if slide_dom_data and slide_dom_data.get("tables"):
      logger.info("[Fallback] Building deterministic native table slide on execution error.")
      t_data = slide_dom_data["tables"][0]
      build_styled_native_table(slide, t_data, Inches(0.8), Inches(1.95), Inches(11.733), Inches(4.5))
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
    else:
      decomp = decompose_slide_image_with_vision(fallback_img_path)
      if decomp and decomp.elements:
        build_native_pptx_slide(slide, decomp, prs, fallback_img_path)
        ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
        ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      else:
        slide.shapes.add_picture(
            fallback_img_path,
            Inches(0),
            Inches(0),
            width=prs.slide_width,
            height=prs.slide_height,
        )


def build_native_pptx_slide(
    slide,
    decomposition: Optional[SlideNativeDecomposition],
    prs: Presentation,
    fallback_img_path: str,
):
  """Builds 100% native editable shapes, text boxes, and tables on a slide, or falls back to picture."""
  if not decomposition or not decomposition.elements:
    slide.shapes.add_picture(
        fallback_img_path,
        Inches(0),
        Inches(0),
        width=prs.slide_width,
        height=prs.slide_height,
    )
    return

  slide_w = prs.slide_width
  slide_h = prs.slide_height

  # 1. Slide Background
  bg_color = _hex_to_pptx_color(
      decomposition.slide_background_hex, default=(255, 255, 255)
  )
  bg_shape = slide.shapes.add_shape(
      MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), slide_w, slide_h
  )
  bg_shape.fill.solid()
  bg_shape.fill.fore_color.rgb = bg_color
  bg_shape.line.fill.background()

  def _in_x(pct: float) -> Inches:
    return Inches(pct / 100.0 * 13.333333)

  def _in_y(pct: float) -> Inches:
    return Inches(pct / 100.0 * 7.5)

  # 2. Containers / Cards (Background layers first)
  for elem in decomposition.elements:
    if elem.element_type == "CONTAINER":
      left = _in_x(elem.left_pct)
      top = _in_y(elem.top_pct)
      width = _in_x(elem.width_pct)
      height = _in_y(elem.height_pct)

      shape_type = (
          MSO_SHAPE.ROUNDED_RECTANGLE
          if (elem.border_radius_pt or 0) > 0
          else MSO_SHAPE.RECTANGLE
      )
      shape = slide.shapes.add_shape(shape_type, left, top, width, height)

      if elem.bg_color_hex:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _hex_to_pptx_color(elem.bg_color_hex)
      else:
        shape.fill.background()

      if elem.border_color_hex:
        shape.line.color.rgb = _hex_to_pptx_color(elem.border_color_hex)
        shape.line.width = Pt(elem.border_width_pt or 1.0)
      else:
        shape.line.fill.background()

      if elem.accent_bar_color_hex:
        bar_w = width - Inches(0.40)
        if bar_w > Inches(0.5):
          bar_shape = slide.shapes.add_shape(
              MSO_SHAPE.ROUNDED_RECTANGLE,
              left + Inches(0.20),
              top + Inches(0.08),
              bar_w,
              Inches(0.06),
          )
          bar_shape.fill.solid()
          bar_shape.fill.fore_color.rgb = _hex_to_pptx_color(elem.accent_bar_color_hex)
          bar_shape.line.fill.background()

  # 3. Dividers
  for elem in decomposition.elements:
    if elem.element_type == "DIVIDER":
      left = _in_x(elem.left_pct)
      top = _in_y(elem.top_pct)
      width = _in_x(elem.width_pct)
      height = _in_y(max(elem.height_pct, 0.2))

      line_shape = slide.shapes.add_shape(
          MSO_SHAPE.RECTANGLE, left, top, width, height
      )
      line_color = _hex_to_pptx_color(elem.border_color_hex or "#E2E8F0")
      line_shape.fill.solid()
      line_shape.fill.fore_color.rgb = line_color
      line_shape.line.fill.background()

  # 4. Sub-cards, Code Blocks, Callouts, Pipeline Steps
  for elem in decomposition.elements:
    if elem.element_type in ("SUB_CARD", "CODE_BLOCK", "CALLOUT_BOX", "PIPELINE_STEP"):
      left = _in_x(elem.left_pct)
      top = _in_y(elem.top_pct)
      width = _in_x(elem.width_pct)
      height = _in_y(elem.height_pct)

      shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
      if elem.bg_color_hex:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _hex_to_pptx_color(elem.bg_color_hex)
      else:
        shape.fill.background()

      if elem.border_color_hex:
        shape.line.color.rgb = _hex_to_pptx_color(elem.border_color_hex)
        shape.line.width = Pt(elem.border_width_pt or 1.0)
      else:
        shape.line.fill.background()

      if elem.text and elem.text.strip():
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.04)
        tf.margin_bottom = Inches(0.04)
        for p_idx, line in enumerate(elem.text.split("\n")):
          p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
          p.text = line
          p.font.name = "Consolas" if elem.element_type == "CODE_BLOCK" else "Noto Sans KR"
          p.font.size = Pt(elem.font_size_pt or (9.0 if elem.element_type == "CODE_BLOCK" else 10.5))
          p.font.bold = bool(elem.is_bold) or (p_idx == 0 and elem.element_type != "CODE_BLOCK")
          if elem.font_color_hex:
            p.font.color.rgb = _hex_to_pptx_color(elem.font_color_hex)

  # 5. Tables
  for elem in decomposition.elements:
    if elem.element_type == "TABLE" and elem.table_matrix:
      matrix = elem.table_matrix
      rows = len(matrix)
      cols = max(len(r) for r in matrix) if rows > 0 else 0
      if rows > 0 and cols > 0:
        left = _in_x(elem.left_pct)
        top = _in_y(elem.top_pct)
        width = _in_x(elem.width_pct)
        height = _in_y(elem.height_pct)

        table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        tbl = table_shape.table
        for r_idx, row in enumerate(matrix):
          for c_idx, cell_val in enumerate(row):
            if c_idx < cols:
              cell = tbl.cell(r_idx, c_idx)
              cell.text = str(cell_val).strip()
              for p in cell.text_frame.paragraphs:
                p.font.name = "Noto Sans KR"
                p.font.size = Pt(11.0 if r_idx > 0 else 12.0)
                p.font.bold = (r_idx == 0)
                p.font.color.rgb = _hex_to_pptx_color(
                    "#1E293B" if r_idx == 0 else "#0F172A"
                )
              if r_idx == 0 and elem.table_header_bg_hex:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _hex_to_pptx_color(
                    elem.table_header_bg_hex
                )

  # 6. Flow Nodes, Flow Arrows, Bridge Connectors, Badges
  for elem in decomposition.elements:
    if elem.element_type in ("FLOW_NODE", "FLOW_ARROW", "BRIDGE_CONNECTOR", "HEADER_BADGE", "ICON_BADGE"):
      left = _in_x(elem.left_pct)
      top = _in_y(elem.top_pct)
      width = _in_x(elem.width_pct)
      height = _in_y(elem.height_pct)

      shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
      if elem.bg_color_hex:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _hex_to_pptx_color(elem.bg_color_hex)
      else:
        shape.fill.background()

      if elem.border_color_hex:
        shape.line.color.rgb = _hex_to_pptx_color(elem.border_color_hex)
        shape.line.width = Pt(elem.border_width_pt or 1.0)
      else:
        shape.line.fill.background()

      if elem.text and elem.text.strip():
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.04)
        tf.margin_right = Inches(0.04)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        p = tf.paragraphs[0]
        p.text = elem.text.strip()
        p.font.name = "Noto Sans KR"
        p.font.size = Pt(elem.font_size_pt or 10.0)
        p.font.bold = bool(elem.is_bold)
        p.font.color.rgb = _hex_to_pptx_color(elem.font_color_hex or "#FFFFFF")
        if elem.alignment == "CENTER":
          p.alignment = PP_ALIGN.CENTER
        elif elem.alignment == "RIGHT":
          p.alignment = PP_ALIGN.RIGHT

  # 7. Text Boxes, Stat Metrics, and Footnotes
  for elem in decomposition.elements:
    if elem.element_type in ("TEXT_BOX", "STAT_METRIC", "FOOTNOTE") and elem.text and elem.text.strip():
      left = _in_x(elem.left_pct)
      top = _in_y(elem.top_pct)
      width = _in_x(elem.width_pct)
      height = _in_y(elem.height_pct)

      txBox = slide.shapes.add_textbox(left, top, width, height)
      tf = txBox.text_frame
      tf.word_wrap = True
      tf.margin_left = Inches(0.04)
      tf.margin_right = Inches(0.04)
      tf.margin_top = Inches(0.02)
      tf.margin_bottom = Inches(0.02)

      lines = [l for l in elem.text.split("\n") if l.strip()]
      for p_idx, line in enumerate(lines):
        p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
        p.text = line.strip()
        p.font.name = "Noto Sans KR"
        p.font.size = Pt(elem.font_size_pt or (9.5 if elem.element_type == "FOOTNOTE" else 14.0))
        p.font.bold = bool(elem.is_bold)
        p.font.color.rgb = _hex_to_pptx_color(
            elem.font_color_hex or ("#64748B" if elem.element_type == "FOOTNOTE" else "#0F172A")
        )
        if elem.alignment == "CENTER":
          p.alignment = PP_ALIGN.CENTER
        elif elem.alignment == "RIGHT":
          p.alignment = PP_ALIGN.RIGHT
        else:
          p.alignment = PP_ALIGN.LEFT


