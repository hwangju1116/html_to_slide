import concurrent.futures
import os
import re
import tempfile
from typing import Any, Dict, List, Optional
import lxml.html
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from app.browser_renderer import (
    FALLBACK_FONTS_DIR,
    FONTS_DIR,
    _capture_slides_cdp,
    _extract_slide_ids,
    _find_chrome_binary,
    _find_free_port,
    _get_font_path,
    _prepare_slide_html,
    _run_async_in_thread,
    _tag_slide_ids_in_html,
    capture_html_slides,
)
from app.color_utils import (
    extract_global_design_tokens,
    extract_root_css_vars,
    hex_to_pptx_color as _hex_to_pptx_color,
    hex_to_rgb_tuple as _hex_to_rgb_tuple,
    parse_color_value,
    resolve_style_color,
)
from app.js_geometry_extractor import JS_SLIDE_GEOMETRY_EXTRACTOR
from app.pptx_native_builders import (
    apply_semantic_styles_to_table,
    audit_and_resolve_slide_collisions,
    build_slide_from_geometry,
    build_styled_native_chart,
    build_styled_native_table,
    ensure_slide_canvas_background,
    ensure_slide_typography_consistency,
    harmonize_deck_presentation_fidelity,
    refine_card_accent_bars,
)
from app.schemas import SlideNativeDecomposition, SlideNativeElement, TextRun
from app.vision_fallback_builder import (
    build_native_pptx_slide,
    build_native_slide_from_code,
    decompose_slide_image_with_vision,
    generate_native_slide_builder_code,
)


def parse_slide_semantic_data(slide_elem, css_vars_raw: Dict[str, str]) -> Dict[str, Any]:
    """Performs deep semantic DOM extraction of slide headers, tables, cell alert colors, badges, and cards."""
    slide_id = slide_elem.get("id", "slide-unknown")
    tag_nodes = slide_elem.xpath(
        './/*[contains(@class, "slide-tag") or contains(@class, "tag") or contains(@class, "badge") or contains(@class, "eyebrow") or contains(@class, "pill")]'
    )
    num_nodes = slide_elem.xpath(
        './/*[contains(@class, "slide-number") or contains(@class, "slide-counter") or contains(@class, "foot__n")]'
    )
    tag = tag_nodes[0].text_content().strip() if tag_nodes else ""
    number = num_nodes[0].text_content().strip() if num_nodes else ""

    h_nodes = slide_elem.xpath(".//h1 | .//h2")
    sub_nodes = slide_elem.xpath(
        './/p[contains(@class, "subtitle") or contains(@class, "premise") or contains(@class, "lead")]'
    )
    title = h_nodes[0].text_content().strip() if h_nodes else ""
    subtitle = sub_nodes[0].text_content().strip() if sub_nodes else ""

    tables_data = []
    for tbl in slide_elem.xpath(".//table"):
        headers = [th.text_content().strip() for th in tbl.xpath(".//thead//th | .//tr[1]/th")]
        rows = []
        body_trs = tbl.xpath(".//tbody//tr")
        if not body_trs:
            all_trs = tbl.xpath(".//tr")
            body_trs = all_trs[1:] if headers and len(all_trs) > 1 else all_trs

        for tr in body_trs:
            row_cells = []
            tr_style = tr.get("style", "")
            tr_color = resolve_style_color(tr_style, css_vars_raw, default=None)
            is_tr_bold = "font-weight: 700" in tr_style or "font-weight: bold" in tr_style

            for c_idx, td in enumerate(tr.xpath("./td | ./th")):
                text = " ".join(td.text_content().split())
                td_style = td.get("style", "")
                cell_color = resolve_style_color(td_style, css_vars_raw, default=tr_color)
                if cell_color is None:
                    inner_styled = td.xpath(".//*[@style]")
                    for st_el in inner_styled:
                        c_col = resolve_style_color(st_el.get("style", ""), css_vars_raw, default=None)
                        if c_col is not None:
                            cell_color = c_col
                            break

                color_rgb = cell_color if cell_color is not None else RGBColor(17, 17, 21)
                is_parent_bold = (
                    is_tr_bold
                    or "font-weight: 700" in td_style
                    or "font-weight: bold" in td_style
                    or (c_idx == 0)
                )
                has_badge = bool(
                    td.xpath('.//*[contains(@class, "math-badge") or contains(@class, "badge")]')
                )
                badge_text = td.xpath(
                    './/*[contains(@class, "math-badge") or contains(@class, "badge")]/text()'
                )
                badge_str = badge_text[0].strip() if badge_text else ""

                child_nodes = td.xpath("child::node()")
                runs = []
                if any(getattr(n, "tag", None) in ("b", "strong", "span") for n in child_nodes):
                    for n in child_nodes:
                        if isinstance(n, str):
                            if n:
                                runs.append({"text": n, "bold": is_parent_bold})
                        elif getattr(n, "tag", None) in ("b", "strong"):
                            runs.append({"text": n.text_content(), "bold": True})
                        else:
                            runs.append({"text": n.text_content(), "bold": is_parent_bold})

                hex_color = f"#{color_rgb[0]:02X}{color_rgb[1]:02X}{color_rgb[2]:02X}"
                row_cells.append({
                    "text": text,
                    "color_hex": hex_color,
                    "color_rgb": (color_rgb[0], color_rgb[1], color_rgb[2]),
                    "is_bold": is_parent_bold or bool(td.xpath(".//b | .//strong")),
                    "has_badge": has_badge,
                    "badge_text": badge_str,
                    "runs": runs,
                })
            if row_cells:
                rows.append(row_cells)

        parent_card = tbl.xpath(
            'ancestor::div[contains(@class, "info-card") or contains(@class, "card") or contains(@class, "panel")]'
        )
        card_title_nodes = (
            parent_card[0].xpath(".//h3/text() | .//h4/text()") if parent_card else []
        )
        card_title = card_title_nodes[0].strip() if card_title_nodes else ""

        tables_data.append({
            "headers": headers,
            "rows": rows,
            "num_cols": max(len(headers), max((len(r) for r in rows), default=0)),
            "num_rows": (1 if headers else 0) + len(rows),
            "in_card": bool(parent_card),
            "card_title": card_title,
        })

    cards_data = []
    card_candidates = slide_elem.xpath(
        './/*[contains(@class, "info-card") or contains(@class, "card") or contains(@class, "panel") or contains(@class, "pillar") or contains(@class, "box") or contains(@class, "runtime") or contains(@class, "attach-card") or contains(@class, "oss-card")]'
    )
    card_elems = [
        c
        for c in card_candidates
        if not any(p != c and c in p.iterdescendants() for p in card_candidates)
    ]
    for c in card_elems:
        if c.xpath(".//table"):
            continue
        c_headings = c.xpath(
            './/h1 | .//h2 | .//h3 | .//h4 | .//*[contains(@class, "title") or contains(@class, "head") or contains(@class, "name")]'
        )
        c_title = c_headings[0].text_content().strip() if c_headings else ""
        paras = [
            p.text_content().strip()
            for p in c.xpath('.//p | .//li | .//*[contains(@class, "f-item")]')
            if p.text_content().strip()
        ]
        badges = [
            b.text_content().strip()
            for b in c.xpath(
                './/*[contains(@class, "badge") or contains(@class, "tag") or contains(@class, "lang")]'
            )
        ]
        cls = c.get("class", "")
        accent = (
            "purple"
            if "purple-accent" in cls
            else (
                "blue"
                if "blue-accent" in cls
                else ("green" if "green-accent" in cls else ("red" if "red-accent" in cls else "none"))
            )
        )
        cards_data.append({
            "title": c_title,
            "paragraphs": paras,
            "badges": badges,
            "accent": accent,
        })

    hl_boxes = []
    for hl in slide_elem.xpath(
        './/*[contains(@class, "highlight-box") or contains(@class, "foot__pocket") or contains(@class, "callout") or contains(@class, "alert")]'
    ):
        txt = hl.text_content().strip()
        if txt:
            hl_boxes.append({
                "title": txt[:40],
                "body": txt,
                "color": "blue",
            })

    math_badges = [
        b.text_content().strip() for b in slide_elem.xpath('.//span[contains(@class, "math-badge")]')
    ]
    raw_html = lxml.html.tostring(slide_elem, encoding="unicode", pretty_print=True)

    return {
        "slide_id": slide_id,
        "tag": tag,
        "number": number,
        "title": title,
        "subtitle": subtitle,
        "tables": tables_data,
        "cards": cards_data,
        "highlight_boxes": hl_boxes,
        "math_badges": math_badges,
        "raw_html": raw_html,
    }


def convert_html_to_pptx(
    html_input: str,
    output_pptx_path: Optional[str] = None,
    target_slide_id: Optional[str] = None,
    base_dir: Optional[str] = None,
    scale_factor: int = 2,
    native_mode: bool = True,
    chrome_binary: Optional[str] = None,
) -> Dict[str, Any]:
    """Dynamically converts any multi-slide HTML file or HTML string to a 16:9 widescreen PowerPoint presentation."""
    if os.path.isfile(html_input):
        html_path = os.path.abspath(html_input)
        if not base_dir:
            base_dir = os.path.dirname(html_path)
        if not output_pptx_path:
            base_name = os.path.splitext(os.path.basename(html_path))[0]
            output_pptx_path = os.path.join(base_dir, f"{base_name}.pptx")
        with open(html_path, "r", encoding="utf-8") as f:
            raw_html_content = f.read()
    else:
        raw_html_content = html_input
        if not base_dir:
            base_dir = os.getcwd()
        if not output_pptx_path:
            output_pptx_path = os.path.join(base_dir, "slide.pptx")

    with tempfile.TemporaryDirectory() as tmp_captures:
        cap_res = capture_html_slides(
            html_input=html_input,
            output_dir=tmp_captures,
            target_slide_id=target_slide_id,
            base_dir=base_dir,
            scale_factor=scale_factor,
            chrome_binary=chrome_binary,
        )
        all_captured_images = cap_res["image_paths"]
        all_captured_geometries = cap_res.get("slide_geometries", [])
        all_slide_ids = cap_res["slide_ids"]
        total_deck_slides = cap_res["slide_count"]

        css_vars_raw = extract_root_css_vars(raw_html_content)
        tagged_html_content, _ = _tag_slide_ids_in_html(raw_html_content)
        try:
            doc = lxml.html.fromstring(tagged_html_content)
        except Exception:
            doc = None

        active_indices = list(range(len(all_slide_ids)))
        if target_slide_id:
            target_clean = str(target_slide_id).strip().lower()
            matched = []
            for idx, sid in enumerate(all_slide_ids):
                sid_clean = sid.lower()
                num_match = re.search(r"\d+", sid_clean)
                num_str = num_match.group(0) if num_match else ""
                if (
                    target_clean == sid_clean
                    or target_clean == f"slide-{sid_clean}"
                    or target_clean == num_str
                    or target_clean == str(idx + 1)
                ):
                    matched.append(idx)
            if matched:
                active_indices = matched

        captured_images = [all_captured_images[i] for i in active_indices]
        captured_geometries = [
            all_captured_geometries[i] if i < len(all_captured_geometries) else None
            for i in active_indices
        ]
        slide_ids = [all_slide_ids[i] for i in active_indices]
        total_slides = len(captured_images)

        slide_dom_list = []
        for sid in slide_ids:
            s_dom = None
            if doc is not None:
                nodes = doc.xpath(f"//*[@id='{sid}']")
                if nodes:
                    s_dom = parse_slide_semantic_data(nodes[0], css_vars_raw)
            slide_dom_list.append(s_dom)

        global_tokens = extract_global_design_tokens(
            first_image_path=captured_images[0] if captured_images else "",
            raw_html=raw_html_content,
            first_slide_geometry=captured_geometries[0] if captured_geometries else None,
        )

        prs = Presentation()
        prs.slide_width = Inches(13.333333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]

        slide_codes = [None] * len(captured_images)
        if native_mode and captured_images:
            indices_needing_codegen = [
                idx
                for idx, _ in enumerate(captured_images)
                if not (
                    captured_geometries[idx]
                    and (
                        captured_geometries[idx].get("cards")
                        or captured_geometries[idx].get("tables")
                        or captured_geometries[idx].get("charts")
                        or captured_geometries[idx].get("title")
                    )
                )
            ]
            if indices_needing_codegen:
                max_workers = min(len(indices_needing_codegen), 2)
                with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                    futures = {
                        executor.submit(
                            generate_native_slide_builder_code,
                            captured_images[idx],
                            slide_dom_data=slide_dom_list[idx],
                            slide_geometry=captured_geometries[idx],
                            design_tokens=global_tokens,
                            slide_index=active_indices[idx] + 1,
                            total_slides=total_deck_slides,
                        ): idx
                        for idx in indices_needing_codegen
                    }
                    for future in concurrent.futures.as_completed(futures):
                        idx = futures[future]
                        try:
                            slide_codes[idx] = future.result()
                        except Exception as e:
                            print(f"[Semantic CodeGen Error for Slide {active_indices[idx]+1}]: {e}")
                            slide_codes[idx] = None

        for idx, img_path in enumerate(captured_images):
            slide = prs.slides.add_slide(blank_layout)
            build_native_slide_from_code(
                slide,
                prs=prs,
                code_str=slide_codes[idx],
                fallback_img_path=img_path,
                design_tokens=global_tokens,
                slide_dom_data=slide_dom_list[idx],
                slide_geometry=captured_geometries[idx],
            )

        harmonize_deck_presentation_fidelity(
            prs,
            default_font=global_tokens.get("font_name", "Pretendard"),
            default_bg_hex=global_tokens.get("bg_color_hex"),
        )

        os.makedirs(os.path.dirname(os.path.abspath(output_pptx_path)), exist_ok=True)
        prs.save(output_pptx_path)

        file_size = os.path.getsize(output_pptx_path)

        return {
            "success": True,
            "slide_count": total_slides,
            "slide_ids": slide_ids,
            "output_pptx_path": output_pptx_path,
            "file_size_bytes": file_size,
        }


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(
        description="Dynamically convert an HTML slide deck to PPTX with Native Editable Vision Decomposition."
    )
    parser.add_argument("input_html", help="Path to the input HTML presentation file")
    parser.add_argument("-o", "--output", help="Path to output PPTX file (optional)")
    parser.add_argument(
        "--scale",
        type=int,
        default=2,
        help="Device scale factor (default: 2 for Retina)",
    )
    parser.add_argument(
        "--native",
        action="store_true",
        default=True,
        help="Use native editable shapes (default: True)",
    )

    args = parser.parse_args()
    result = convert_html_to_pptx(
        args.input_html,
        output_pptx_path=args.output,
        scale_factor=args.scale,
        native_mode=args.native,
    )
    print(f"[SUCCESS] Converted {result['slide_count']} slides to PPTX:")
    print(f"File: {result['output_pptx_path']} ({result['file_size_bytes']:,} bytes)")
