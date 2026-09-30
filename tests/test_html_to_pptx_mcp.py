import asyncio
import unittest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.html_to_pptx_converter import (
    _extract_slide_ids,
    _find_chrome_binary,
    audit_and_resolve_slide_collisions,
    build_native_slide_from_code,
    ensure_slide_typography_consistency,
    extract_global_design_tokens,
)
from app.mcp_server import mcp


class TestHtmlToPptxMcp(unittest.TestCase):
    def test_mcp_tools_registered(self):
        tools = asyncio.run(mcp.list_tools())
        tool_names = [t.name for t in tools]
        self.assertIn("convert_html_to_pptx_mcp", tool_names)
        self.assertIn("capture_html_slides_mcp", tool_names)
        self.assertIn("recreate_editable_slide_mcp", tool_names)

    def test_extract_slide_ids(self):
        html_multi = """
        <div id="slide-1" class="slide">Slide 1</div>
        <div id="slide-2" class="slide">Slide 2</div>
        """
        self.assertEqual(_extract_slide_ids(html_multi), ["slide-1", "slide-2"])

    def test_find_chrome_binary(self):
        self.assertTrue(bool(_find_chrome_binary()))

    def test_collision_auditor_deduplicates_backgrounds(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Add two full-bleed background shapes
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        self.assertEqual(len(slide.shapes), 2)

        audit_and_resolve_slide_collisions(slide)
        self.assertEqual(len(slide.shapes), 1)

    def test_ensure_slide_typography_consistency(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
        tf = shape.text_frame
        p1 = tf.paragraphs[0]
        p1.text = "Heading"
        ensure_slide_typography_consistency(slide, default_font="Pretendard")
        self.assertEqual(p1.font.name, "Pretendard")

    def test_extract_global_design_tokens(self):
        html_sample = "<style>:root { --brand-color: #2563eb; --text-main: #111115; }</style>"
        tokens = extract_global_design_tokens(raw_html=html_sample)
        self.assertEqual(tokens["brand_color_hex"], "#2563eb")
        self.assertEqual(tokens["text_main_hex"], "#111115")

    def test_resolve_gcp_project_and_remote_helpers(self):
        import json
        import os
        import tempfile
        from app.genai_client import resolve_gcp_project
        from app.mcp_server import (
            _is_raw_html,
            _stage_for_download,
            convert_html_to_pptx_mcp,
        )

        self.assertEqual(resolve_gcp_project("custom-customer-proj"), "custom-customer-proj")
        self.assertTrue(_is_raw_html("<!DOCTYPE html><html><body>Slide</body></html>"))
        self.assertTrue(_is_raw_html('<div class="slide">Hello</div>'))
        self.assertFalse(_is_raw_html("/nonexistent/path/slide.html"))

        err_res = json.loads(convert_html_to_pptx_mcp("/nonexistent/local/path/slide.html"))
        self.assertEqual(err_res["status"], "error")
        self.assertIn("raw HTML markup", err_res["error"])

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            tmp.write(b"dummy-pptx")
            tmp_path = tmp.name
        try:
            os.environ["PUBLIC_BASE_URL"] = "https://html-to-pptx-mcp-123.run.app"
            staged_name, url = _stage_for_download(tmp_path)
            self.assertTrue(staged_name.endswith(".pptx"))
            self.assertEqual(url, f"https://html-to-pptx-mcp-123.run.app/downloads/{staged_name}")
        finally:
            os.environ.pop("PUBLIC_BASE_URL", None)
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_dark_theme_background_preserved_in_harmonization(self):
        from pptx.dml.color import RGBColor
        from app.html_to_pptx_converter import harmonize_deck_presentation_fidelity

        dark_html = "<style>:root { --bg-dark: #0B0E14; --lg-red-bright: #FD3153; }</style>"
        tokens = extract_global_design_tokens(
            raw_html=dark_html,
            first_slide_geometry={"slideBgColor": "rgb(13, 17, 25)"},
        )
        self.assertEqual(tokens["bg_color_hex"], "#0D1119")
        self.assertEqual(tokens["text_main_hex"], "#F8FAFC")

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(13, 17, 25)

        harmonize_deck_presentation_fidelity(prs, default_font="Pretendard", default_bg_hex="#0D1119")
        self.assertEqual(bg.fill.fore_color.rgb, RGBColor(13, 17, 25))

    def test_build_styled_native_chart_types(self):
        from pptx.enum.chart import XL_CHART_TYPE
        from app.html_to_pptx_converter import build_styled_native_chart

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        line_chart = build_styled_native_chart(
            slide,
            {
                "rect": {"left": 7.5, "top": 2.0, "width": 5.0, "height": 2.5},
                "chartConfig": {
                    "type": "line",
                    "labels": ["2022", "2023", "2024", "2025", "2026E"],
                    "datasets": [{"label": "AI 탑재율 (%)", "data": [18, 29, 44, 59, 74.2], "borderColor": "#FD3153", "tension": 0.35}],
                    "legendDisplay": True,
                },
            },
            is_dark=True,
            slide_bg_rgb=(13, 17, 25),
        )
        self.assertTrue(line_chart.has_chart)
        self.assertEqual(line_chart.chart.chart_type, XL_CHART_TYPE.LINE_MARKERS)

        bar_chart = build_styled_native_chart(
            slide,
            {
                "rect": {"left": 0.5, "top": 2.0, "width": 6.5, "height": 3.5},
                "chartConfig": {
                    "type": "bar",
                    "indexAxis": "y",
                    "labels": ["중국", "미국", "인도"],
                    "datasets": [{"label": "GDP (PPP)", "data": [35.2, 28.8, 14.6], "backgroundColor": ["#64748B", "#38BDF8", "#10B981"]}],
                    "legendDisplay": False,
                },
            },
            is_dark=True,
            slide_bg_rgb=(13, 17, 25),
        )
        self.assertTrue(bar_chart.has_chart)
        self.assertEqual(bar_chart.chart.chart_type, XL_CHART_TYPE.BAR_CLUSTERED)

        doughnut_chart = build_styled_native_chart(
            slide,
            {
                "rect": {"left": 0.8, "top": 2.2, "width": 4.0, "height": 3.0},
                "chartConfig": {
                    "type": "doughnut",
                    "labels": ["단발성 판매", "가전 구독", "소모품/AI"],
                    "datasets": [{"data": [62, 28, 10], "backgroundColor": ["#334155", "#FD3153", "#10B981"]}],
                    "cutout": "68%",
                },
            },
            is_dark=True,
            slide_bg_rgb=(13, 17, 25),
        )
        self.assertTrue(doughnut_chart.has_chart)
        self.assertEqual(doughnut_chart.chart.chart_type, XL_CHART_TYPE.DOUGHNUT)

    def test_alpha_blended_rgba_color_parsing(self):
        from pptx.dml.color import RGBColor
        from app.html_to_pptx_converter import parse_color_value

        # Glass card rgba(255, 255, 255, 0.05) blended against dark background (13, 17, 25)
        blended = parse_color_value("rgba(255, 255, 255, 0.05)", default=RGBColor(255, 255, 255), bg_blend_rgb=(13, 17, 25))
        self.assertLess(blended[0], 40)
        self.assertLess(blended[1], 40)
        self.assertLess(blended[2], 45)

    def test_resolve_style_color_ignores_background_and_border_color(self):
        from pptx.dml.color import RGBColor
        from app.color_utils import resolve_style_color

        # Should not falsely match background-color or border-color when extracting text color
        self.assertIsNone(
            resolve_style_color("background-color: #ff0000; border-color: #00ff00;", {}, default=None)
        )
        # Should match standalone color property
        self.assertEqual(
            resolve_style_color("background-color: #ff0000; color: #2563eb;", {}, default=None),
            RGBColor(37, 99, 235),
        )

    def test_prepare_slide_html_preserves_hyphenated_active_classes(self):
        from app.browser_renderer import _prepare_slide_html

        raw = """
        <div id="slide-1" class="slide active is-active">
          <button class="tab is-active nav-active">Tab</button>
        </div>
        <div id="slide-2" class="slide">
          <span class="badge-active">Slide 2</span>
        </div>
        """
        prepared = _prepare_slide_html(raw, "slide-2", 1, "/tmp")
        self.assertIn("is-active", prepared)
        self.assertIn("nav-active", prepared)
        self.assertIn("badge-active", prepared)
        self.assertIn('id="slide-2" class="slide active"', prepared)


if __name__ == "__main__":
    unittest.main()


