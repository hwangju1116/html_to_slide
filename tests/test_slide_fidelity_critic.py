import unittest
from unittest.mock import patch, MagicMock
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from app.slide_fidelity_critic import (
    extract_slide_text_summary,
    run_code_on_test_slide,
    audit_slide_fidelity,
)
from app.schemas import SlideFidelityReport


class TestSlideFidelityCritic(unittest.TestCase):
    def test_extract_slide_text_summary(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
        tf = shape.text_frame
        tf.paragraphs[0].text = "Core Strategy Title"

        summary = extract_slide_text_summary(slide)
        self.assertIn("Core Strategy Title", summary)

    def test_run_code_on_test_slide_valid(self):
        code = """
def build_slide(prs, slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(5), Inches(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    tf = shape.text_frame
    tf.paragraphs[0].text = "Valid Test Slide"
"""
        success, slide, prs, err = run_code_on_test_slide(code)
        self.assertTrue(success)
        self.assertIsNotNone(slide)
        self.assertEqual(err, "")
        self.assertEqual(len(slide.shapes), 1)

    def test_run_code_on_test_slide_syntax_error(self):
        code = "def build_slide(prs, slide):\n    this is invalid python syntax !!!"
        success, slide, prs, err = run_code_on_test_slide(code)
        self.assertFalse(success)
        self.assertIn("syntax", err.lower())

    def test_audit_empty_slide_fails_rule_check(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # empty slide
        report = audit_slide_fidelity(
            original_image_input=b"dummy_bytes",
            generated_code="def build_slide(prs, slide): pass",
            slide=slide,
        )
        self.assertFalse(report.passed)
        self.assertLess(report.score, 7)
        self.assertIn("empty", report.missing_or_truncated_text[0].lower())

    @patch("app.slide_fidelity_critic.get_genai_client")
    def test_audit_slide_fidelity_mock_pass(self, mock_get_client):
        mock_client = MagicMock()
        mock_resp = MagicMock()
        mock_resp.text = '{"passed": true, "score": 9, "zero_omission_verified": true, "missing_or_truncated_text": [], "design_and_layout_issues": [], "corrective_instructions": ""}'
        mock_client.models.generate_content.return_value = mock_resp
        mock_get_client.return_value = mock_client

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        s1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13), Inches(7))
        s2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(5), Inches(3))
        s2.text_frame.paragraphs[0].text = "Full Content Present with 100% Zero Omission"

        report = audit_slide_fidelity(
            original_image_input=b"fake_image_bytes",
            generated_code="def build_slide(prs, slide): pass",
            slide=slide,
        )
        self.assertTrue(report.passed)
        self.assertEqual(report.score, 9)
        self.assertTrue(report.zero_omission_verified)


if __name__ == "__main__":
    unittest.main()
