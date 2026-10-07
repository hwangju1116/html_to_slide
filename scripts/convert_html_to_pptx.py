import argparse
import json
import logging
import os
from pathlib import Path
import re
import sys
from typing import Any, Dict, List, Optional
import lxml.html
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

SKILL_ROOT = str(Path(__file__).resolve().parent.parent)
if SKILL_ROOT not in sys.path:
    sys.path.insert(0, SKILL_ROOT)

from scripts.browser_renderer import (
    FALLBACK_FONTS_DIR,
    FONTS_DIR,
    _extract_slide_ids,
    _extract_slides_cdp,
    _find_chrome_binary,
    _find_free_port,
    _get_font_path,
    _prepare_slide_html,
    _run_async_in_thread,
    _tag_slide_ids_in_html,
    extract_html_slides_geometry,
)
from scripts.color_utils import (
    extract_global_design_tokens,
    extract_root_css_vars,
    hex_to_pptx_color as _hex_to_pptx_color,
    hex_to_rgb_tuple as _hex_to_rgb_tuple,
    parse_color_value,
    resolve_style_color,
)
from scripts.js_geometry_extractor import JS_SLIDE_GEOMETRY_EXTRACTOR
from scripts.pptx_native_builders import (
    apply_semantic_styles_to_table,
    audit_and_resolve_slide_collisions,
    build_native_slide_from_code,
    build_slide_from_geometry,
    build_styled_native_chart,
    build_styled_native_table,
    ensure_slide_canvas_background,
    ensure_slide_typography_consistency,
    harmonize_deck_presentation_fidelity,
    refine_card_accent_bars,
)

logger = logging.getLogger(__name__)


def parse_slide_semantic_data(slide_elem, css_vars_raw: Dict[str, str]) -> Dict[str, Any]:
    slide_id = slide_elem.get("id", "slide-unknown")
    tag_nodes = slide_elem.xpath(
        './/*[contains(@class, "slide-tag") or contains(@class, "tag") or contains(@class, "badge") or contains(@class, "eyebrow") or contains(@class, "pill")]'
    )
    num_nodes = slide_elem.xpath(
        './/*[contains(@class, "slide-number") or contains(@class, "slide-counter") or contains(@class, "foot__n")]'
    )
    tag = tag_nodes[0].text_content().strip() if tag_nodes else ""
    number = num_nodes[0].text_content().strip() if num_nodes else ""

    h_nodes = slide_elem.xpath(".//h1 | .//h2")
    sub_nodes = slide_elem.xpath(
        './/p[contains(@class, "subtitle") or contains(@class, "premise") or contains(@class, "lead")]'
    )
    title = h_nodes[0].text_content().strip() if h_nodes else ""
    subtitle = sub_nodes[0].text_content().strip() if sub_nodes else ""

    tables_data = []
    for tbl in slide_elem.xpath(".//table"):
        headers = [th.text_content().strip() for th in tbl.xpath(".//thead//th | .//tr[1]/th")]
        rows = []
        body_trs = tbl.xpath(".//tbody//tr")
        if not body_trs:
            all_trs = tbl.xpath(".//tr")
            body_trs = all_trs[1:] if headers and len(all_trs) > 1 else all_trs

        for tr in body_trs:
            row_cells = []
            tr_style = tr.get("style", "")
            tr_color = resolve_style_color(tr_style, css_vars_raw, default=None)
            is_tr_bold = "font-weight: 700" in tr_style or "font-weight: bold" in tr_style

            for c_idx, td in enumerate(tr.xpath("./td | ./th")):
                text = " ".join(td.text_content().split())
                td_style = td.get("style", "")
                cell_color = resolve_style_color(td_style, css_vars_raw, default=tr_color)
                if cell_color is None:
                    inner_styled = td.xpath(".//*[@style]")
                    for st_el in inner_styled:
                        c_col = resolve_style_color(st_el.get("style", ""), css_vars_raw, default=None)
                        if c_col is not None:
                            cell_color = c_col
                            break

                color_rgb = cell_color if cell_color is not None else RGBColor(17, 17, 21)
                is_parent_bold = (
                    is_tr_bold
                    or "font-weight: 700" in td_style
                    or "font-weight: bold" in td_style
                    or (c_idx == 0)
                )
                has_badge = bool(
                    td.xpath('.//*[contains(@class, "math-badge") or contains(@class, "badge")]')
                )
                badge_text = td.xpath(
                    './/*[contains(@class, "math-badge") or contains(@class, "badge")]/text()'
                )
                badge_str = badge_text[0].strip() if badge_text else ""

                child_nodes = td.xpath("child::node()")
                runs = []
                if any(getattr(n, "tag", None) in ("b", "strong", "span") for n in child_nodes):
                    for n in child_nodes:
                        if isinstance(n, str):
                            if n:
                                runs.append({"text": n, "bold": is_parent_bold})
                        elif getattr(n, "tag", None) in ("b", "strong"):
                            runs.append({"text": n.text_content(), "bold": True})
                        else:
                            runs.append({"text": n.text_content(), "bold": is_parent_bold})

                hex_color = f"{chr(35)}{color_rgb[0]:02X}{color_rgb[1]:02X}{color_rgb[2]:02X}"
                row_cells.append({
                    "text": text,
                    "color_hex": hex_color,
                    "color_rgb": (color_rgb[0], color_rgb[1], color_rgb[2]),
                    "is_bold": is_parent_bold,
                    "has_badge": has_badge,
                    "badge_text": badge_str,
                    "runs": runs,
                })
            if row_cells:
                rows.append(row_cells)

        parent_card = tbl.xpath(
            'ancestor::div[contains(@class, "info-card") or contains(@class, "card") or contains(@class, "panel")]'
        )
        card_title_nodes = (
            parent_card[0].xpath(".//h3/text() | .//h4/text()") if parent_card else []
        )
        card_title = card_title_nodes[0].strip() if card_title_nodes else ""

        tables_data.append({
            "headers": headers,
            "rows": rows,
            "num_cols": max(len(headers), max((len(r) for r in rows), default=0)),
            "num_rows": (1 if headers else 0) + len(rows),
            "in_card": bool(parent_card),
            "card_title": card_title,
        })

    cards_data = []
    card_candidates = slide_elem.xpath(
        './/*[contains(@class, "info-card") or contains(@class, "card") or contains(@class, "panel") or contains(@class, "pillar") or contains(@class, "box") or contains(@class, "runtime") or contains(@class, "attach-card") or contains(@class, "oss-card")]'
    )
    card_elems = [
        c
        for c in card_candidates
        if not any(p != c and c in p.iterdescendants() for p in card_candidates)
    ]
    for c in card_elems:
        if c.xpath(".//table"):
            continue
        c_headings = c.xpath(
            './/h1 | .//h2 | .//h3 | .//h4 | .//*[contains(@class, "title") or contains(@class, "head") or contains(@class, "name")]'
        )
        c_title = c_headings[0].text_content().strip() if c_headings else ""
        paras = [
            p.text_content().strip()
            for p in c.xpath('.//p | .//li | .//*[contains(@class, "f-item")]')
            if p.text_content().strip()
        ]
        badges = [
            b.text_content().strip()
            for b in c.xpath(
                './/*[contains(@class, "badge") or contains(@class, "tag") or contains(@class, "lang")]'
            )
        ]
        cls = c.get("class", "")
        accent = (
            "purple"
            if "purple-accent" in cls
            else (
                "blue"
                if "blue-accent" in cls
                else ("green" if "green-accent" in cls else ("red" if "red-accent" in cls else "none"))
            )
        )
        cards_data.append({
            "title": c_title,
            "paragraphs": paras,
            "badges": badges,
            "accent": accent,
        })

    hl_boxes = []
    for hl in slide_elem.xpath(
        './/*[contains(@class, "highlight-box") or contains(@class, "foot__pocket") or contains(@class, "callout") or contains(@class, "alert")]'
    ):
        txt = hl.text_content().strip()
        if txt:
            hl_boxes.append({
                "title": txt[:40],
                "body": txt,
                "color": "blue",
            })

    math_badges = [
        b.text_content().strip() for b in slide_elem.xpath('.//span[contains(@class, "math-badge")]')
    ]
    raw_html = lxml.html.tostring(slide_elem, encoding="unicode", pretty_print=True)

    return {
        "slide_id": slide_id,
        "tag": tag,
        "number": number,
        "title": title,
        "subtitle": subtitle,
        "tables": tables_data,
        "cards": cards_data,
        "highlight_boxes": hl_boxes,
        "math_badges": math_badges,
        "raw_html": raw_html,
    }


def format_slide_manifest_markdown(
    slide_idx: int,
    slide_id: str,
    screenshot_path: Optional[str],
    design_tokens: Dict[str, Any],
    slide_dom: Optional[Dict[str, Any]],
    slide_geom: Optional[Dict[str, Any]],
) -> str:
    font_family = design_tokens.get("font_name", "Pretendard")
    brand_hex = design_tokens.get("brand_color_hex", f"{chr(35)}2563EB")
    bg_hex = (slide_geom or {}).get("slideBgColor") or design_tokens.get("bg_color_hex", f"{chr(35)}FFFFFF")
    text_main_hex = design_tokens.get("text_main_hex", f"{chr(35)}111115")

    lines = [
        f"================================================================================",
        f"{chr(35)}{chr(35)} Slide {slide_idx} (`{slide_id}`) — `def build_slide_{slide_idx}(prs, slide):`",
        f"================================================================================",
        f"- **Rendered 16:9 Screenshot**: `{screenshot_path or 'N/A'}`",
        f"- **Canvas Background Color**: `{bg_hex}` (gradient=`{(slide_geom or {}).get('slideBgGradient')}`)",
        f"- **Brand Primary Color**: `{brand_hex}`",
        f"- **Primary Font Family**: `{font_family}`",
        f"- **Primary Text Color**: `{text_main_hex}`",
    ]

    if slide_geom:
        if slide_geom.get("headerBadges"):
            for hb_i, hb in enumerate(slide_geom["headerBadges"], 1):
                hbr = hb.get("rect", {})
                hbs = hb.get("styles", {})
                lines.append(
                    f"- **Header Badge {hb_i}**: `{hb.get('text')}` at `left={hbr.get('left')}\", top={hbr.get('top')}\", width={hbr.get('width')}\", height={hbr.get('height')}\"` "
                    f"(bg=`{hbs.get('backgroundColor')}`, color=`{hbs.get('color')}`, fontSize=`{hbs.get('fontSizePt')}pt`)"
                )
        if slide_geom.get("num"):
            nm = slide_geom["num"]
            nmr = nm.get("rect", {})
            lines.append(
                f"- **Slide Number**: `{nm.get('text')}` at `left={nmr.get('left')}\", top={nmr.get('top')}\", width={nmr.get('width')}\", height={nmr.get('height')}\"`"
            )
        if slide_geom.get("title"):
            t = slide_geom["title"]
            tr = t.get("rect", {})
            ts = t.get("styles", {})
            lines.append(
                f"- **Main Title**: `{t.get('text')}` at `left={tr.get('left')}\", top={tr.get('top')}\", width={tr.get('width')}\", height={tr.get('height')}\"` "
                f"(`fontSizePt={ts.get('fontSizePt')}`, `color={ts.get('color')}`, `runs={t.get('runs')}`)"
            )
        if slide_geom.get("sub"):
            sb = slide_geom["sub"]
            sbr = sb.get("rect", {})
            sbs = sb.get("styles", {})
            lines.append(
                f"- **Subtitle / Premise**: `{sb.get('text')}` at `left={sbr.get('left')}\", top={sbr.get('top')}\", width={sbr.get('width')}\", height={sbr.get('height')}\"` "
                f"(`fontSizePt={sbs.get('fontSizePt')}`, `color={sbs.get('color')}`, `runs={sb.get('runs')}`)"
            )
        if slide_geom.get("desc"):
            for d_i, d in enumerate(slide_geom["desc"], 1):
                dr = d.get("rect", {})
                ds = d.get("styles", {})
                lines.append(
                    f"- **Description {d_i}**: `{d.get('text')}` at `left={dr.get('left')}\", top={dr.get('top')}\", width={dr.get('width')}\", height={dr.get('height')}\"` (`color={ds.get('color')}`)"
                )
        if slide_geom.get("decorShapes"):
            lines.append(f"- **Decorative / Pseudo-Element Shapes ({len(slide_geom['decorShapes'])})**: `{slide_geom['decorShapes']}`")
        if slide_geom.get("images"):
            lines.append(
                f"- **Images / Vector SVGs ({len(slide_geom['images'])})**: "
                + ", ".join(
                    f"`{im.get('tagName')}` at `{im.get('rect')}` (svgPrimitives={len(im.get('svgPrimitives') or [])})"
                    for im in slide_geom["images"]
                )
            )

        if slide_geom.get("tables"):
            lines.append(f"- **Native Tables ({len(slide_geom['tables'])})**:")
            for t_i, tbl in enumerate(slide_geom["tables"], 1):
                tr = tbl.get("rect", {})
                lines.append(
                    f"  - Table {t_i}: `{tbl.get('num_rows')} rows x {tbl.get('num_cols')} cols` at "
                    f"`left={tr.get('left')}\", top={tr.get('top')}\", width={tr.get('width')}\", height={tr.get('height')}\"` "
                    f"(`colRatios={tbl.get('colRatios')}`, `inCard={tbl.get('inCard')}`, `hasShadow={tbl.get('hasShadow')}`, `cardTitle={tbl.get('cardTitle')}`)"
                )
                lines.append(f"    - Headers: `{tbl.get('headers', [])}`")
                for r_i, row in enumerate(tbl.get("rows", []), 1):
                    r_desc = [
                        f"{c['text']} (color={c.get('color')}, bg={c.get('bgColor')}, bold={c.get('is_bold')}, runs={c.get('runs')})"
                        for c in row
                    ]
                    lines.append(f"    - Row {r_i}: {' | '.join(r_desc)}")
            lines.append(
                "  - **[TABLE BUILDER INSTRUCTIONS]**: Use `build_styled_native_table(slide, tbl, Inches(left), Inches(top), Inches(width), Inches(height), font_name=...)` or `build_slide_from_geometry(slide, slide_geometry)` so exact column ratios, cell colors, inline badges/tags (`.score`, `.tag`), `<code>` highlights, `.sub` multi-line spans, and partial `<b>` bolding are preserved."
            )

        if slide_geom.get("charts"):
            lines.append(f"- **Native Charts ({len(slide_geom['charts'])})**:")
            for ch_i, ch in enumerate(slide_geom["charts"], 1):
                chr_r = ch.get("rect", {})
                cfg = ch.get("chartConfig") or {}
                lines.append(
                    f"  - Chart {ch_i} (`id={ch.get('id')}`): type=`{cfg.get('type')}`, indexAxis=`{cfg.get('indexAxis')}` at "
                    f"`left={chr_r.get('left')}\", top={chr_r.get('top')}\", width={chr_r.get('width')}\", height={chr_r.get('height')}\"`"
                )
                lines.append(f"    - Labels: `{cfg.get('labels')}`")
                lines.append(f"    - Datasets: `{cfg.get('datasets')}`")
            lines.append(
                "  - **[CHART BUILDER INSTRUCTIONS]**: Call `build_styled_native_chart(slide, ch, font_name=..., is_dark=..., slide_bg_rgb=...)` for each item in `charts_data` (or `slide_geometry['charts']`) to render 100% native editable PowerPoint charts with exact dataset colors."
            )

        if slide_geom.get("cards"):
            lines.append(f"- **Container Cards & Panels ({len(slide_geom['cards'])})**:")
            for c_i, c in enumerate(slide_geom["cards"], 1):
                cr = c.get("rect", {})
                cs = c.get("styles", {})
                lines.append(
                    f"  - Card {c_i}: title=`{c.get('title')}` at `left={cr.get('left')}\", top={cr.get('top')}\", width={cr.get('width')}\", height={cr.get('height')}\"` "
                    f"(fill=`{cs.get('backgroundColor')}`, gradient=`{cs.get('gradient')}`, hasShadow=`{cs.get('hasShadow')}`, border=`{cs.get('borderColor')}`, "
                    f"topAccent=`{c.get('topAccentColor') if c.get('hasTopAccent') else 'none'}`, "
                    f"leftAccent=`{c.get('leftAccentColor') if c.get('hasLeftAccent') else 'none'}`)"
                )
                if c.get("tag"):
                    lines.append(f"    - Tag: `{c['tag'].get('text')}` (color=`{c['tag'].get('styles', {}).get('color')}`, bg=`{c['tag'].get('styles', {}).get('backgroundColor')}`)")
                if c.get("titleStyles"):
                    lines.append(f"    - Title Style: color=`{c['titleStyles'].get('color')}`, size=`{c['titleStyles'].get('fontSizePt')}pt`, runs=`{c.get('titleRuns')}`")
                if c.get("desc"):
                    lines.append(f"    - Card Desc: `{c['desc'].get('text')}` (color=`{c['desc'].get('styles', {}).get('color')}`)")
                if c.get("isBridge") and c.get("bridge"):
                    lines.append(f"    - Bridge Badge: `{c['bridge']}`")
                if c.get("isAttachCard") and c.get("attachData"):
                    lines.append(f"    - Attach Card Data: `{c['attachData']}`")
                if c.get("pipelineSteps"):
                    lines.append(f"    - Pipeline / Process Steps ({len(c['pipelineSteps'])}):")
                    for ps_i, ps in enumerate(c["pipelineSteps"], 1):
                        psr = ps.get("rect", {})
                        pss = ps.get("styles", {})
                        lines.append(
                            f"      - Step {ps_i} (`stepNum={ps.get('stepNum')}`): title=`{ps.get('title')}` (color=`{(ps.get('titleStyles') or {}).get('color')}`), "
                            f"desc=`{ps.get('desc')}` (color=`{(ps.get('descStyles') or {}).get('color')}`) at "
                            f"`left={psr.get('left')}\", top={psr.get('top')}\", width={psr.get('width')}\", height={psr.get('height')}\"` "
                            f"(bg=`{pss.get('backgroundColor')}`, border=`{pss.get('borderColor')}`, align=`{pss.get('textAlign')}`)"
                        )
                if c.get("flows"):
                    lines.append(f"    - Process Flows ({len(c['flows'])}):")
                    for fl_i, fl in enumerate(c["flows"], 1):
                        flr = fl.get("rect", {})
                        fls = fl.get("styles", {})
                        lines.append(
                            f"      - Flow {fl_i} at `left={flr.get('left')}\", top={flr.get('top')}\", width={flr.get('width')}\", height={flr.get('height')}\"` (bg=`{fls.get('backgroundColor')}`):"
                        )
                        for n_i, nd in enumerate(fl.get("nodes", []), 1):
                            ndr = nd.get("rect", {})
                            nds = nd.get("styles", {})
                            lines.append(
                                f"        - Node {n_i}: `{nd.get('text')}` (`isArrow={nd.get('isArrow')}`, `isHighlight={nd.get('isHighlight')}`, "
                                f"bg=`{nds.get('backgroundColor')}`, color=`{nds.get('color')}`, border=`{nds.get('borderColor')}`) at "
                                f"`left={ndr.get('left')}\", top={ndr.get('top')}\", width={ndr.get('width')}\", height={ndr.get('height')}\"`"
                            )
                if c.get("fItems"):
                    lines.append(f"    - Feature Items ({len(c['fItems'])}):")
                    for fi_i, fi in enumerate(c["fItems"], 1):
                        fir = fi.get("rect", {})
                        lines.append(
                            f"      - FItem {fi_i}: icon=`{fi.get('icon')}` (bg=`{(fi.get('iconStyles') or {}).get('backgroundColor')}`, color=`{(fi.get('iconStyles') or {}).get('color')}`), "
                            f"text=`{fi.get('text')}` (color=`{(fi.get('textStyles') or {}).get('color')}`) at "
                            f"`left={fir.get('left')}\", top={fir.get('top')}\", width={fir.get('width')}\", height={fir.get('height')}\"`"
                        )
                if c.get("subCards"):
                    lines.append(f"    - Sub-Cards ({len(c['subCards'])}):")
                    for sc_i, sc in enumerate(c["subCards"], 1):
                        scr = sc.get("rect", {})
                        scs = sc.get("styles", {})
                        lines.append(
                            f"      - SubCard {sc_i}: head=`{sc.get('head')}` (badge=`{sc.get('badge')}`, headColor=`{(sc.get('headStyles') or {}).get('color')}`), "
                            f"items=`{sc.get('items')}` at `left={scr.get('left')}\", top={scr.get('top')}\", width={scr.get('width')}\", height={scr.get('height')}\"` "
                            f"(bg=`{scs.get('backgroundColor')}`, gradient=`{scs.get('gradient')}`, border=`{scs.get('borderColor')}`, textColor=`{scs.get('color')}`)"
                        )
                if c.get("badges"):
                    for b_i, b in enumerate(c["badges"], 1):
                        br = b.get("rect", {})
                        bs = b.get("styles", {})
                        lines.append(
                            f"    - Badge {b_i}: `{b.get('text')}` at `left={br.get('left')}\", top={br.get('top')}\", width={br.get('width')}\", height={br.get('height')}\"` (bg=`{bs.get('backgroundColor')}`, color=`{bs.get('color')}`)"
                        )
                if c.get("codeBlocks"):
                    for cb_i, cb in enumerate(c["codeBlocks"], 1):
                        cbr = cb.get("rect", {})
                        lines.append(
                            f"    - CodeBlock {cb_i}: `{cb.get('text')}` at `left={cbr.get('left')}\", top={cbr.get('top')}\", width={cbr.get('width')}\", height={cbr.get('height')}\"` (runs=`{cb.get('runs')}`)"
                        )
                if c.get("paragraphs"):
                    for p_i, p in enumerate(c["paragraphs"], 1):
                        pr = p.get("rect", {})
                        ps = p.get("styles", {})
                        lines.append(
                            f"    - Paragraph {p_i} (`isList={p.get('isList')}`): `{p.get('text')}` at "
                            f"`left={pr.get('left')}\", top={pr.get('top')}\", width={pr.get('width')}\", height={pr.get('height')}\"` "
                            f"(color=`{ps.get('color')}`, bold=`{ps.get('isBold')}`, size=`{ps.get('fontSizePt')}pt`, runs=`{p.get('runs')}`)"
                        )

        if slide_geom.get("highlightBoxes"):
            for h_i, hl in enumerate(slide_geom["highlightBoxes"], 1):
                hr = hl.get("rect", {})
                hs = hl.get("styles", {})
                lines.append(
                    f"- **Highlight / Callout Banner {h_i}**: `{hl.get('text')}` at `left={hr.get('left')}\", top={hr.get('top')}\", width={hr.get('width')}\", height={hr.get('height')}\"` "
                    f"(bg=`{hs.get('backgroundColor')}`, border=`{hs.get('borderColor')}`, color=`{hs.get('color')}`, runs=`{hl.get('runs')}`)"
                )

        if slide_geom.get("footnotes"):
            for fn_i, fn in enumerate(slide_geom["footnotes"], 1):
                fnr = fn.get("rect", {})
                fns = fn.get("styles", {})
                lines.append(
                    f"- **Footnote {fn_i}**: `{fn.get('text')}` at `left={fnr.get('left')}\", top={fnr.get('top')}\", width={fnr.get('width')}\", height={fnr.get('height')}\"` (`color={fns.get('color')}`, `fontSizePt={fns.get('fontSizePt')}`)"
                )

        geom_for_prompt = dict(slide_geom)
        if geom_for_prompt.get("charts"):
            clean_charts = []
            for ch in geom_for_prompt["charts"]:
                ch_copy = dict(ch)
                if "dataUrl" in ch_copy:
                    ch_copy["dataUrl"] = "<omitted_base64_png>"
                clean_charts.append(ch_copy)
            geom_for_prompt["charts"] = clean_charts

        lines.extend([
            "",
            f"{chr(35)}{chr(35)}{chr(35)} [EXACT MEASURED BROWSER LAYOUT GEOMETRY (PIXEL-PERFECT INCHES)]",
            "```json",
            json.dumps(geom_for_prompt, ensure_ascii=False, indent=2),
            "```",
        ])

    if slide_dom and slide_dom.get("raw_html"):
        raw_html_snippet = slide_dom["raw_html"][:6000]
        lines.extend([
            "",
            f"{chr(35)}{chr(35)}{chr(35)} [SLIDE ORIGINAL HTML SOURCE]",
            "```html",
            raw_html_snippet,
            "```",
        ])

    lines.append("")
    return "\n".join(lines)


def prepare_html_for_host_ai(
    html_input: str,
    prepare_dir: str,
    target_slide_id: Optional[str] = None,
    base_dir: Optional[str] = None,
    scale_factor: int = 2,
    chrome_binary: Optional[str] = None,
) -> Dict[str, Any]:
    prepare_dir = os.path.abspath(prepare_dir)
    os.makedirs(prepare_dir, exist_ok=True)

    if os.path.isfile(html_input):
        html_path = os.path.abspath(html_input)
        if not base_dir:
            base_dir = os.path.dirname(html_path)
        with open(html_path, "r", encoding="utf-8") as f:
            raw_html_content = f.read()
    else:
        raw_html_content = html_input
        if not base_dir:
            base_dir = os.getcwd()

    ext_res = extract_html_slides_geometry(
        html_input=html_input,
        target_slide_id=target_slide_id,
        base_dir=base_dir,
        scale_factor=scale_factor,
        chrome_binary=chrome_binary,
        capture_all_screenshots=True,
        screenshot_dir=prepare_dir,
    )
    slide_geometries = ext_res.get("slide_geometries", [])
    screenshot_paths = ext_res.get("screenshot_paths", [])
    slide_ids = ext_res["slide_ids"]

    css_vars_raw = extract_root_css_vars(raw_html_content)
    tagged_html_content, _ = _tag_slide_ids_in_html(raw_html_content)
    try:
        doc = lxml.html.fromstring(tagged_html_content)
    except Exception:
        doc = None

    slide_dom_list = []
    for sid in slide_ids:
        s_dom = None
        if doc is not None:
            nodes = doc.xpath(f"//*[@id='{sid}']")
            if nodes:
                s_dom = parse_slide_semantic_data(nodes[0], css_vars_raw)
        slide_dom_list.append(s_dom)

    global_tokens = extract_global_design_tokens(
        raw_html=raw_html_content,
        first_slide_geometry=slide_geometries[0] if slide_geometries else None,
    )

    geom_json_path = os.path.join(prepare_dir, "geometry.json")
    with open(geom_json_path, "w", encoding="utf-8") as f_geom:
        json.dump(
            {
                "design_tokens": global_tokens,
                "slide_ids": slide_ids,
                "slide_geometries": slide_geometries,
            },
            f_geom,
            ensure_ascii=False,
            indent=2,
        )

    manifest_sections = [
        f"{chr(35)} [AUTHENTIC SLIDE STRUCTURE & STYLE MANIFEST]",
        "",
        f"{chr(35)}{chr(35)} [STRICT INSTRUCTIONS FOR THE SLIDE BUILDER AI]",
        "1. **Visual & Coordinate Grounding**: Inspect each slide's rendered 16:9 PNG screenshot (`slide-<N>.png`) alongside its `[EXACT MEASURED BROWSER LAYOUT GEOMETRY]` below. Use the exact `left`, `top`, `width`, `height` inch coordinates (`Inches(...)`) and measured `fontSizePt` from the browser geometry so there are zero overlaps, zero layout shifts, and 100% consistent header/footer typography across all slides.",
        "2. **Exact Colors, Gradients, Shadows, & Inline Highlights**: Preserve 100% of the exact RGB/HEX colors (`color`, `backgroundColor`, `borderColor`, `topAccentColor`, `leftAccentColor`, colored text `runs`). Use `apply_linear_gradient_fill(shape, grad_info)` for CSS linear-gradient bars/cards (never split gradients into multiple solid rectangles), `apply_drop_shadow(shape)` for `hasShadow: true` cards, and `apply_run_highlight(run, rgb_color)` or `render_rich_runs_into_text_frame(tf, runs, ...)` for inline `<code>` (`Consolas` + `#eef0f5` highlight) and inline `.score`/`.tag` badges.",
        "3. **Processes, Flows, Sub-Cards, & Vector SVGs**: Render every `pipelineSteps` item, `flows` container/node/arrow, `fItems` icon/text row, `subCards` box, and `decorShapes` item as native PowerPoint shapes at their exact measured coordinates. Pure geometric `<svg>` illustrations (`svgPrimitives`) are automatically rendered as native editable vector shapes by `build_svg_primitives` / `embed_extracted_slide_images`.",
        "4. **Native Charts & Tables**: Always use `build_styled_native_chart(slide, ch, ...)` for `charts_data` and `build_styled_native_table(slide, tbl, ...)` for `table_data`. Calling `build_slide_from_geometry(slide, slide_geometry)` automatically renders all tables (with `colRatios`, rounded card wrapper, drop shadow, partial `<b>` bolding, `.sub` lines, `<code>` highlights, and inline `.score`/`.tag` badges), charts, cards, gradients, and vector SVGs.",
        "5. **16:9 Canvas**: Assume `prs.slide_width = Inches(13.333333)` and `prs.slide_height = Inches(7.5)`.",
        "",
    ]
    for idx, sid in enumerate(slide_ids):
        shot_p = screenshot_paths[idx] if idx < len(screenshot_paths) else None
        s_dom = slide_dom_list[idx] if idx < len(slide_dom_list) else None
        s_geom = slide_geometries[idx] if idx < len(slide_geometries) else None
        manifest_sections.append(
            format_slide_manifest_markdown(idx + 1, sid, shot_p, global_tokens, s_dom, s_geom)
        )

    manifest_path = os.path.join(prepare_dir, "manifest.md")
    with open(manifest_path, "w", encoding="utf-8") as f_man:
        f_man.write("\n".join(manifest_sections))

    return {
        "status": "prepared",
        "success": True,
        "slide_count": len(slide_ids),
        "slide_ids": slide_ids,
        "prepare_dir": prepare_dir,
        "screenshot_paths": screenshot_paths,
        "manifest_path": manifest_path,
        "geometry_json_path": geom_json_path,
    }


def convert_html_to_pptx(
    html_input: str,
    output_pptx_path: Optional[str] = None,
    target_slide_id: Optional[str] = None,
    base_dir: Optional[str] = None,
    scale_factor: int = 2,
    chrome_binary: Optional[str] = None,
    custom_builder_code: Optional[str] = None,
) -> Dict[str, Any]:
    if os.path.isfile(html_input):
        html_path = os.path.abspath(html_input)
        if not base_dir:
            base_dir = os.path.dirname(html_path)
        if not output_pptx_path:
            base_name = os.path.splitext(os.path.basename(html_path))[0]
            output_pptx_path = os.path.join(base_dir, f"{base_name}.pptx")
        with open(html_path, "r", encoding="utf-8") as f:
            raw_html_content = f.read()
    else:
        raw_html_content = html_input
        if not base_dir:
            base_dir = os.getcwd()
        if not output_pptx_path:
            output_pptx_path = os.path.join(base_dir, "slide.pptx")

    ext_res = extract_html_slides_geometry(
        html_input=html_input,
        target_slide_id=target_slide_id,
        base_dir=base_dir,
        scale_factor=scale_factor,
        chrome_binary=chrome_binary,
    )
    slide_geometries = ext_res.get("slide_geometries", [])
    slide_image_buffers = ext_res.get("slide_image_buffers", [])
    slide_ids = ext_res["slide_ids"]
    total_slides = len(slide_ids)

    css_vars_raw = extract_root_css_vars(raw_html_content)
    tagged_html_content, _ = _tag_slide_ids_in_html(raw_html_content)
    try:
        doc = lxml.html.fromstring(tagged_html_content)
    except Exception:
        doc = None

    slide_dom_list = []
    for sid in slide_ids:
        s_dom = None
        if doc is not None:
            nodes = doc.xpath(f"//*[@id='{sid}']")
            if nodes:
                s_dom = parse_slide_semantic_data(nodes[0], css_vars_raw)
        slide_dom_list.append(s_dom)

    global_tokens = extract_global_design_tokens(
        raw_html=raw_html_content,
        first_slide_geometry=slide_geometries[0] if slide_geometries else None,
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for idx in range(total_slides):
        slide = prs.slides.add_slide(blank_layout)
        build_native_slide_from_code(
            slide,
            prs=prs,
            code_str=custom_builder_code,
            screenshot_source=slide_image_buffers[idx] if idx < len(slide_image_buffers) else None,
            design_tokens=global_tokens,
            slide_dom_data=slide_dom_list[idx],
            slide_geometry=slide_geometries[idx] if idx < len(slide_geometries) else None,
            slide_index=idx + 1,
        )

    harmonize_deck_presentation_fidelity(
        prs,
        default_font=global_tokens.get("font_name", "Pretendard"),
        default_bg_hex=global_tokens.get("bg_color_hex"),
    )

    output_pptx_path = os.path.abspath(output_pptx_path)
    os.makedirs(os.path.dirname(output_pptx_path), exist_ok=True)
    prs.save(output_pptx_path)

    file_size = os.path.getsize(output_pptx_path)

    return {
        "status": "success",
        "success": True,
        "slide_count": total_slides,
        "slide_ids": slide_ids,
        "output_pptx_path": output_pptx_path,
        "file_size_bytes": file_size,
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert single- or multi-slide HTML presentations into 100%% native editable 16:9 PowerPoint (.pptx) decks."
    )
    parser.add_argument(
        "input_html",
        help="Path to the input .html presentation file (or raw HTML markup string).",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Output .pptx file path (defaults to <input_stem>.pptx in the same directory).",
    )
    parser.add_argument(
        "-s",
        "--slide",
        dest="target_slide_id",
        help="Optional 1-based slide index or HTML id (e.g. '1' or 'slide-3') to convert only a single slide.",
    )
    parser.add_argument(
        "--scale",
        type=int,
        default=2,
        help="Device scale factor for Chromium rendering (default: 2 for Retina).",
    )
    parser.add_argument(
        "--prepare",
        dest="prepare_dir",
        help="Render 16:9 PNG screenshots and generate manifest.md + geometry.json in the specified directory for Host AI Vision code generation.",
    )
    parser.add_argument(
        "--builder-script",
        help="Optional path to a custom Python script defining `build_slide(prs, slide)` or `build_slide_<N>(prs, slide)` generated by the Host AI.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Emit structured JSON output to stdout.",
    )

    args = parser.parse_args()

    if args.prepare_dir:
        prep_res = prepare_html_for_host_ai(
            html_input=args.input_html,
            prepare_dir=args.prepare_dir,
            target_slide_id=args.target_slide_id,
            scale_factor=args.scale,
        )
        if args.json:
            print(json.dumps(prep_res, ensure_ascii=False, indent=2))
        else:
            print(f"[PREPARED] {prep_res['slide_count']} slide(s) in {prep_res['prepare_dir']}")
            print(f"Manifest: {prep_res['manifest_path']}")
        return

    custom_code = None
    if args.builder_script and os.path.isfile(args.builder_script):
        with open(args.builder_script, "r", encoding="utf-8") as f_code:
            custom_code = f_code.read()

    result = convert_html_to_pptx(
        html_input=args.input_html,
        output_pptx_path=args.output,
        target_slide_id=args.target_slide_id,
        scale_factor=args.scale,
        custom_builder_code=custom_code,
    )

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(f"[SUCCESS] Converted {result['slide_count']} slide(s) to PPTX:")
        print(f"File: {result['output_pptx_path']} ({result['file_size_bytes']:,} bytes)")


if __name__ == "__main__":
    main()
