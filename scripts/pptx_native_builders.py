import base64
import inspect
import io
import logging
import os
from pathlib import Path
import re
import sys
from typing import Any, Dict, List, Optional, Union
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt

SKILL_ROOT = str(Path(__file__).resolve().parent.parent)
if SKILL_ROOT not in sys.path:
  sys.path.insert(0, SKILL_ROOT)

from scripts.color_utils import (
    hex_to_pptx_color as _hex_to_pptx_color,
    hex_to_rgb_tuple as _hex_to_rgb_tuple,
    parse_color_value,
)

logger = logging.getLogger(__name__)


def _disable_default_pptx_table_style(tbl) -> None:
  try:
    tbl_pr = tbl._tbl.tblPr
    if tbl_pr is not None:
      tbl_pr.attrib["firstRow"] = "0"
      tbl_pr.attrib["bandRow"] = "0"
      tbl_pr.attrib["firstCol"] = "0"
      tbl_pr.attrib["lastCol"] = "0"
      tbl_pr.attrib["bandCol"] = "0"
  except Exception:
    pass


def _apply_cell_bottom_border(cell, border_rgb: RGBColor, width_emu: int = 9525) -> None:
  try:
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
      for existing in tc_pr.xpath(f"./{tag}"):
        tc_pr.remove(existing)
    hex_val = f"{border_rgb[0]:02X}{border_rgb[1]:02X}{border_rgb[2]:02X}"
    ln_l = parse_xml(f"<a:lnL {nsdecls('a')}><a:noFill/></a:lnL>")
    ln_r = parse_xml(f"<a:lnR {nsdecls('a')}><a:noFill/></a:lnR>")
    ln_t = parse_xml(f"<a:lnT {nsdecls('a')}><a:noFill/></a:lnT>")
    ln_b = parse_xml(
        f'<a:lnB {nsdecls("a")} w="{width_emu}">'
        f'<a:solidFill><a:srgbClr val="{hex_val}"/></a:solidFill>'
        f"</a:lnB>"
    )
    tc_pr.append(ln_l)
    tc_pr.append(ln_r)
    tc_pr.append(ln_t)
    tc_pr.append(ln_b)
  except Exception:
    pass


def _style_paragraph_and_runs(
    p,
    font_name: str,
    font_size_pt: float,
    bold: bool,
    color_rgb: RGBColor,
) -> None:
  p.font.name = font_name
  p.font.size = Pt(font_size_pt)
  p.font.bold = bold
  p.font.color.rgb = color_rgb
  for r in p.runs:
    r.font.name = font_name
    r.font.size = Pt(font_size_pt)
    r.font.bold = bold
    r.font.color.rgb = color_rgb


def build_styled_native_table(
    slide,
    table_data: Dict[str, Any],
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
    card_border_rgb: Optional[RGBColor] = None,
    font_name: str = "Pretendard",
    col_ratios: Optional[List[float]] = None,
):
  is_dark = table_data.get("is_dark", False)
  slide_bg = table_data.get("slide_bg_rgb", (15, 23, 42) if is_dark else (255, 255, 255))
  blend_rgb = (slide_bg[0], slide_bg[1], slide_bg[2])

  if card_border_rgb is None:
    card_border_rgb = RGBColor(51, 65, 85) if is_dark else RGBColor(233, 219, 237)

  if left < Inches(0.35):
    left = Inches(0.35)
  max_avail_w = Inches(13.333) - left - Inches(0.35)
  if width > max_avail_w:
    width = max_avail_w

  skip_card_wrapper = bool(table_data.get("skip_card_wrapper", False))
  in_card = bool(table_data.get("in_card") or table_data.get("inCard"))
  card_rect = table_data.get("cardRect")
  card_styles = table_data.get("cardStyles") or {}
  card_title = (table_data.get("cardTitle") or table_data.get("card_title") or "").strip() if not skip_card_wrapper else ""
  card_title_rect = table_data.get("cardTitleRect")
  card_title_styles = table_data.get("cardTitleStyles") or {}

  if not skip_card_wrapper and in_card:
    if card_rect and isinstance(card_rect, dict):
      c_left = Inches(card_rect.get("left", left.inches))
      c_top = Inches(card_rect.get("top", top.inches))
      c_w = Inches(card_rect.get("width", width.inches))
      c_h = Inches(card_rect.get("height", height.inches))
    else:
      c_left, c_top, c_w, c_h = left, top, width, height

    card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, c_w, c_h)
    card_shape.fill.solid()
    card_bg = parse_color_value(
        card_styles.get("backgroundColor"),
        table_data.get("card_bg_rgb") or (RGBColor(30, 41, 59) if is_dark else RGBColor(255, 255, 255)),
        bg_blend_rgb=blend_rgb,
    )
    card_shape.fill.fore_color.rgb = card_bg
    c_border = parse_color_value(
        card_styles.get("borderColor"),
        card_border_rgb,
        bg_blend_rgb=blend_rgb,
    )
    card_shape.line.color.rgb = c_border
    card_shape.line.width = Pt(1.0)

    if card_title:
      if card_title_rect and isinstance(card_title_rect, dict):
        ct_l = Inches(card_title_rect.get("left", c_left.inches + 0.20))
        ct_t = Inches(card_title_rect.get("top", c_top.inches + 0.14))
        ct_w = Inches(max(2.0, card_title_rect.get("width", c_w.inches - 0.40)))
        ct_h = Inches(max(0.32, card_title_rect.get("height", 0.36)))
      else:
        ct_l = c_left + Inches(0.20)
        ct_t = c_top + Inches(0.14)
        ct_w = c_w - Inches(0.40)
        ct_h = Inches(0.36)
      tx = slide.shapes.add_textbox(ct_l, ct_t, ct_w, ct_h)
      tf_ct = tx.text_frame
      tf_ct.word_wrap = True
      p_ct = tf_ct.paragraphs[0]
      p_ct.text = card_title
      ct_color = parse_color_value(
          card_title_styles.get("color"),
          RGBColor(147, 197, 253) if is_dark else RGBColor(95, 0, 128),
      )
      ct_size = float(card_title_styles.get("fontSizePt") or 13.5)
      _style_paragraph_and_runs(p_ct, font_name, ct_size, True, ct_color)

  if skip_card_wrapper or (in_card and card_rect):
    t_top = top
    t_height = height
    t_left = left
    t_width = width
  elif in_card:
    t_top = top + (Inches(0.52) if card_title else Inches(0.12))
    t_height = height - (Inches(0.64) if card_title else Inches(0.24))
    t_left = left + Inches(0.12)
    t_width = width - Inches(0.24)
  else:
    t_top = top
    t_height = height
    t_left = left
    t_width = width

  num_rows = table_data["num_rows"]
  num_cols = table_data["num_cols"]
  table_shape = slide.shapes.add_table(num_rows, num_cols, t_left, t_top, t_width, t_height)
  tbl = table_shape.table
  _disable_default_pptx_table_style(tbl)

  if col_ratios and len(col_ratios) == num_cols:
    total_r = sum(col_ratios)
    for c_idx, r in enumerate(col_ratios):
      tbl.columns[c_idx].width = int(t_width * (r / total_r))
  else:
    eq_w = int(t_width / num_cols)
    for c_idx in range(num_cols):
      tbl.columns[c_idx].width = eq_w

  default_border_rgb = RGBColor(51, 65, 85) if is_dark else RGBColor(233, 224, 238)

  if table_data.get("headers"):
    for c_idx, h in enumerate(table_data["headers"]):
      if c_idx >= num_cols:
        break
      h_text = h["text"] if isinstance(h, dict) else str(h)
      h_bg = h.get("bgColor") if isinstance(h, dict) else None
      h_col = h.get("color") if isinstance(h, dict) else None
      h_size = float(h.get("fontSizePt") or 11.0) if isinstance(h, dict) else 11.0
      h_border_str = h.get("borderColor") if isinstance(h, dict) else None

      if h_bg:
        header_bg_rgb = parse_color_value(
            h_bg,
            default=RGBColor(15, 23, 42) if is_dark else RGBColor(244, 237, 246),
            bg_blend_rgb=blend_rgb,
        )
      elif is_dark:
        header_bg_rgb = RGBColor(15, 23, 42)
      else:
        header_bg_rgb = RGBColor(244, 237, 246)

      if h_col:
        header_fg_rgb = parse_color_value(
            h_col,
            default=RGBColor(148, 163, 184) if is_dark else RGBColor(95, 0, 128),
        )
      elif is_dark:
        header_fg_rgb = RGBColor(148, 163, 184)
      else:
        header_fg_rgb = RGBColor(95, 0, 128)

      cell = tbl.cell(0, c_idx)
      cell.text = h_text
      cell.fill.solid()
      cell.fill.fore_color.rgb = header_bg_rgb
      cell.vertical_anchor = MSO_ANCHOR.MIDDLE
      cell.margin_left = Inches(0.14)
      cell.margin_right = Inches(0.14)
      cell.margin_top = Inches(0.08)
      cell.margin_bottom = Inches(0.08)
      b_rgb = parse_color_value(h_border_str, default=default_border_rgb, bg_blend_rgb=blend_rgb)
      _apply_cell_bottom_border(cell, b_rgb, width_emu=12700)
      p = cell.text_frame.paragraphs[0]
      _style_paragraph_and_runs(p, font_name, h_size, True, header_fg_rgb)

  for r_idx, row in enumerate(table_data["rows"]):
    row_num = r_idx + (1 if table_data.get("headers") else 0)
    if row_num >= num_rows:
      break
    for c_idx, cell_data in enumerate(row):
      if c_idx >= num_cols:
        break
      cell = tbl.cell(row_num, c_idx)
      cell.vertical_anchor = MSO_ANCHOR.MIDDLE
      cell.margin_left = Inches(0.14)
      cell.margin_right = Inches(0.14)
      cell.margin_top = Inches(0.08)
      cell.margin_bottom = Inches(0.08)
      cell.fill.solid()

      cell_bg = cell_data.get("bgColor")
      if cell_bg:
        row_bg_rgb = parse_color_value(
            cell_bg,
            default=RGBColor(30, 41, 59) if is_dark else RGBColor(255, 255, 255),
            bg_blend_rgb=blend_rgb,
        )
      elif is_dark:
        row_bg_rgb = RGBColor(30, 41, 59) if r_idx % 2 == 0 else RGBColor(22, 30, 46)
      else:
        row_bg_rgb = RGBColor(255, 255, 255)
      cell.fill.fore_color.rgb = row_bg_rgb

      c_border_rgb = parse_color_value(
          cell_data.get("borderColor"),
          default=default_border_rgb,
          bg_blend_rgb=blend_rgb,
      )
      _apply_cell_bottom_border(cell, c_border_rgb, width_emu=9525)

      tf = cell.text_frame
      tf.word_wrap = True
      p = tf.paragraphs[0]
      cell_font_pt = float(cell_data.get("fontSizePt") or 10.5)

      default_cell_fg = RGBColor(241, 245, 249) if is_dark else RGBColor(17, 17, 21)
      if "color" in cell_data and cell_data["color"]:
        col_rgb = parse_color_value(cell_data["color"], default=default_cell_fg)
      elif "color_rgb" in cell_data and cell_data["color_rgb"]:
        col_rgb = RGBColor(*cell_data["color_rgb"])
      else:
        col_rgb = default_cell_fg

      bg_lum = row_bg_rgb[0] * 0.299 + row_bg_rgb[1] * 0.587 + row_bg_rgb[2] * 0.114
      fg_lum = col_rgb[0] * 0.299 + col_rgb[1] * 0.587 + col_rgb[2] * 0.114
      if bg_lum < 120 and fg_lum < 90:
        col_rgb = RGBColor(241, 245, 249)
      elif bg_lum >= 150 and fg_lum > 180:
        col_rgb = RGBColor(17, 17, 21)

      if cell_data.get("lines") and len(cell_data["lines"]) > 0:
        for l_idx, line_info in enumerate(cell_data["lines"]):
          lp = tf.paragraphs[0] if l_idx == 0 else tf.add_paragraph()
          lp.text = line_info.get("text", "")
          l_size = float(line_info.get("fontSizePt") or (11.0 if l_idx == 0 else 9.5))
          l_bold = bool(line_info.get("isBold", l_idx == 0 and cell_data.get("is_bold")))
          l_col = parse_color_value(line_info.get("color"), default=col_rgb)
          l_lum = l_col[0] * 0.299 + l_col[1] * 0.587 + l_col[2] * 0.114
          if bg_lum < 120 and l_lum < 75:
            l_col = RGBColor(203, 213, 225)
          elif bg_lum >= 150 and l_lum > 190:
            l_col = RGBColor(51, 65, 85)
          _style_paragraph_and_runs(lp, font_name, l_size, l_bold, l_col)
      elif cell_data.get("has_badge"):
        p.text = cell_data.get("badge_text") or cell_data["text"]
        b_fg = parse_color_value(
            cell_data.get("badge_color") or cell_data.get("color"),
            default=RGBColor(96, 165, 250) if is_dark else RGBColor(95, 0, 128),
        )
        _style_paragraph_and_runs(p, font_name, cell_font_pt, True, b_fg)
      elif cell_data.get("runs") and len(cell_data["runs"]) > 1:
        p.text = ""
        for r_info in cell_data["runs"]:
          run = p.add_run()
          run.text = r_info["text"]
          run.font.name = font_name
          run.font.size = Pt(cell_font_pt)
          run.font.bold = bool(r_info.get("bold"))
          r_col = parse_color_value(r_info.get("color"), default=col_rgb)
          run.font.color.rgb = r_col
      else:
        p.text = cell_data["text"]
        _style_paragraph_and_runs(p, font_name, cell_font_pt, bool(cell_data.get("is_bold")), col_rgb)

  return table_shape


def build_styled_native_chart(
    slide,
    chart_data: Dict[str, Any],
    font_name: str = "Pretendard",
    is_dark: bool = False,
    slide_bg_rgb: tuple[int, int, int] = (15, 23, 42),
):
  r = chart_data.get("rect", {})
  left = Inches(max(0.2, float(r.get("left", 0.8))))
  top = Inches(max(0.2, float(r.get("top", 1.8))))
  width = Inches(max(1.2, float(r.get("width", 5.0))))
  height = Inches(max(1.0, float(r.get("height", 3.2))))

  cfg = chart_data.get("chartConfig")
  blend_rgb = (slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2])

  if cfg and isinstance(cfg, dict):
    try:
      c_type = str(cfg.get("type", "bar")).lower()
      options = cfg.get("options") if isinstance(cfg.get("options"), dict) else {}
      index_axis = str(cfg.get("indexAxis") or options.get("indexAxis", "x")).lower()

      xl_type = None
      if c_type == "line":
        xl_type = XL_CHART_TYPE.LINE_MARKERS
      elif c_type == "bar" and index_axis == "y":
        xl_type = XL_CHART_TYPE.BAR_CLUSTERED
      elif c_type == "bar":
        xl_type = XL_CHART_TYPE.COLUMN_CLUSTERED
      elif c_type == "doughnut":
        xl_type = XL_CHART_TYPE.DOUGHNUT
      elif c_type == "pie":
        xl_type = XL_CHART_TYPE.PIE
      elif c_type == "radar":
        xl_type = XL_CHART_TYPE.RADAR

      data_dict = cfg.get("data") if isinstance(cfg.get("data"), dict) else {}
      raw_labels = cfg.get("labels") if isinstance(cfg.get("labels"), list) else (data_dict.get("labels") or [])
      raw_datasets = cfg.get("datasets") if isinstance(cfg.get("datasets"), list) else (data_dict.get("datasets") or [])

      if xl_type is not None and len(raw_labels) > 0 and len(raw_datasets) > 0:
        is_horiz_bar = xl_type == XL_CHART_TYPE.BAR_CLUSTERED
        labels = [str(lbl) for lbl in (reversed(raw_labels) if is_horiz_bar else raw_labels)]

        chart_data_obj = CategoryChartData()
        chart_data_obj.categories = labels

        for ds in raw_datasets:
          raw_vals = ds.get("data") or []
          vals_iter = reversed(raw_vals) if is_horiz_bar else raw_vals
          vals = []
          for v in vals_iter:
            try:
              vals.append(float(v))
            except Exception:
              vals.append(0.0)
          chart_data_obj.add_series(str(ds.get("label", "")), tuple(vals))

        chart_shape = slide.shapes.add_chart(xl_type, left, top, width, height, chart_data_obj)
        chart = chart_shape.chart

        try:
          cs = chart._element
          for existing_spPr in cs.xpath("./c:spPr"):
            cs.remove(existing_spPr)
          spPr_cs = parse_xml(f"<c:spPr {nsdecls('c', 'a')}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>")
          chart_el_list = cs.xpath("./c:chart")
          if chart_el_list:
            cs.insert(cs.index(chart_el_list[0]) + 1, spPr_cs)
          else:
            cs.append(spPr_cs)

          pa_list = cs.xpath("./c:chart/c:plotArea")
          if pa_list:
            pa = pa_list[0]
            for existing_pa_spPr in pa.xpath("./c:spPr"):
              pa.remove(existing_pa_spPr)
            spPr_pa = parse_xml(f"<c:spPr {nsdecls('c', 'a')}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>")
            pa.append(spPr_pa)
        except Exception:
          pass

        if xl_type == XL_CHART_TYPE.DOUGHNUT:
          try:
            hole_pct = 65
            cutout_val = cfg.get("cutout") or options.get("cutout")
            if cutout_val:
              m_cut = re.search(r"\d+", str(cutout_val))
              if m_cut:
                hole_pct = max(10, min(90, int(m_cut.group(0))))
            dc_list = chart._element.xpath("./c:chart/c:plotArea/c:doughnutChart")
            if dc_list:
              dc = dc_list[0]
              for hs in dc.xpath("./c:holeSize"):
                dc.remove(hs)
              dc.append(parse_xml(f'<c:holeSize {nsdecls("c")} val="{hole_pct}"/>'))
          except Exception:
            pass

        plugins_cfg = options.get("plugins") if isinstance(options.get("plugins"), dict) else {}
        legend_cfg = plugins_cfg.get("legend") if isinstance(plugins_cfg.get("legend"), dict) else {}
        default_show_legend = len(raw_datasets) > 1 or xl_type in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE)
        if "legendDisplay" in cfg and cfg["legendDisplay"] is not None:
          show_legend = bool(cfg["legendDisplay"])
        else:
          show_legend = bool(legend_cfg.get("display", default_show_legend))
        if show_legend:
          chart.has_legend = True
          pos_str = str(cfg.get("legendPosition") or legend_cfg.get("position", "top")).lower()
          if pos_str == "bottom":
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
          elif pos_str == "right":
            chart.legend.position = XL_LEGEND_POSITION.RIGHT
          elif pos_str == "left":
            chart.legend.position = XL_LEGEND_POSITION.LEFT
          else:
            chart.legend.position = XL_LEGEND_POSITION.TOP
          chart.legend.include_in_layout = False
          chart.legend.font.name = font_name
          chart.legend.font.size = Pt(9.5)
          chart.legend.font.color.rgb = RGBColor(203, 213, 225) if is_dark else RGBColor(51, 65, 85)
        else:
          chart.has_legend = False

        if xl_type not in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE, XL_CHART_TYPE.RADAR):
          axis_fg = RGBColor(148, 163, 184) if is_dark else RGBColor(100, 116, 139)
          grid_hex = "1E293B" if is_dark else "E2E8F0"
          scales_cfg = options.get("scales") if isinstance(options.get("scales"), dict) else {}

          try:
            cat_ax = chart.category_axis
            cat_ax.tick_labels.font.name = font_name
            cat_ax.tick_labels.font.size = Pt(9.0)
            cat_ax.tick_labels.font.color.rgb = axis_fg
            cat_ax.has_major_gridlines = False
            for sp in cat_ax._element.xpath("./c:spPr"):
              cat_ax._element.remove(sp)
            cat_ax._element.append(
                parse_xml(f'<c:spPr {nsdecls("c", "a")}><a:ln w="9525"><a:solidFill><a:srgbClr val="{grid_hex}"/></a:solidFill></a:ln></c:spPr>')
            )
          except Exception:
            pass

          try:
            val_ax = chart.value_axis
            val_scale_key = "x" if is_horiz_bar else "y"
            val_scale_cfg = scales_cfg.get(val_scale_key) if isinstance(scales_cfg.get(val_scale_key), dict) else {}
            explicit_display = cfg.get(f"{val_scale_key}Display")
            if explicit_display is False or val_scale_cfg.get("display") is False:
              val_ax.visible = False
              val_ax.has_major_gridlines = False
            else:
              val_ax.tick_labels.font.name = font_name
              val_ax.tick_labels.font.size = Pt(9.0)
              val_ax.tick_labels.font.color.rgb = axis_fg
              val_ax.has_major_gridlines = True
              mg = val_ax.major_gridlines
              for sp in mg._element.xpath("./c:spPr"):
                mg._element.remove(sp)
              mg._element.append(
                  parse_xml(f'<c:spPr {nsdecls("c", "a")}><a:ln w="9525"><a:solidFill><a:srgbClr val="{grid_hex}"/></a:solidFill></a:ln></c:spPr>')
              )
              for sp in val_ax._element.xpath("./c:spPr"):
                val_ax._element.remove(sp)
              val_ax._element.append(
                  parse_xml(f'<c:spPr {nsdecls("c", "a")}><a:ln><a:noFill/></a:ln></c:spPr>')
              )
          except Exception:
            pass

        default_palette = [
            RGBColor(225, 29, 72),
            RGBColor(59, 130, 246),
            RGBColor(16, 185, 129),
            RGBColor(245, 158, 11),
            RGBColor(139, 92, 246),
        ]
        for s_idx, series in enumerate(chart.series):
          if s_idx >= len(raw_datasets):
            break
          ds = raw_datasets[s_idx]
          fallback_col = default_palette[s_idx % len(default_palette)]
          bg_col = ds.get("backgroundColor")
          border_col = ds.get("borderColor")

          if xl_type == XL_CHART_TYPE.LINE_MARKERS:
            line_col_str = border_col if isinstance(border_col, str) else (bg_col if isinstance(bg_col, str) else None)
            line_rgb = parse_color_value(line_col_str, default=fallback_col, bg_blend_rgb=blend_rgb)
            series.format.line.color.rgb = line_rgb
            series.format.line.width = Pt(float(ds.get("borderWidth", 2.5)))
            if ds.get("borderDash"):
              series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if float(ds.get("tension", 0) or 0) > 0:
              try:
                for sm in series._element.xpath("./c:smooth"):
                  series._element.remove(sm)
                series._element.append(parse_xml(f'<c:smooth {nsdecls("c")} val="1"/>'))
              except Exception:
                pass
            try:
              hex_col = f"{line_rgb[0]:02X}{line_rgb[1]:02X}{line_rgb[2]:02X}"
              for mk in series._element.xpath("./c:marker"):
                series._element.remove(mk)
              mk_xml = parse_xml(
                  f'<c:marker {nsdecls("c", "a")}>'
                  f'<c:symbol val="circle"/><c:size val="6"/>'
                  f'<c:spPr><a:solidFill><a:srgbClr val="{hex_col}"/></a:solidFill>'
                  f'<a:ln w="12700"><a:solidFill><a:srgbClr val="{hex_col}"/></a:solidFill></a:ln></c:spPr>'
                  f"</c:marker>"
              )
              series._element.insert(2, mk_xml)
            except Exception:
              pass
          else:
            if isinstance(bg_col, list) and len(bg_col) > 0:
              col_list = list(reversed(bg_col)) if is_horiz_bar else list(bg_col)
              for pt_idx, pt in enumerate(series.points):
                c_str = col_list[pt_idx % len(col_list)]
                pt_rgb = parse_color_value(c_str, default=fallback_col, bg_blend_rgb=blend_rgb)
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = pt_rgb
                if xl_type in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE):
                  pt.format.line.color.rgb = RGBColor(*blend_rgb)
                  pt.format.line.width = Pt(1.5)
                else:
                  pt.format.line.fill.background()
            else:
              c_str = bg_col if isinstance(bg_col, str) else (border_col if isinstance(border_col, str) else None)
              s_rgb = parse_color_value(c_str, default=fallback_col, bg_blend_rgb=blend_rgb)
              series.format.fill.solid()
              series.format.fill.fore_color.rgb = s_rgb
              series.format.line.fill.background()

        return chart_shape
    except Exception as chart_err:
      logger.warning("[Native Chart Builder Warning]: %s", chart_err)

  data_url = chart_data.get("dataUrl")
  if data_url and isinstance(data_url, str) and "base64," in data_url:
    try:
      b64_part = data_url.split("base64,", 1)[1]
      img_bytes = base64.b64decode(b64_part)
      return slide.shapes.add_picture(io.BytesIO(img_bytes), left, top, width=width, height=height)
    except Exception as pic_err:
      logger.warning("[Chart Picture Fallback Warning]: %s", pic_err)
  return None


def apply_semantic_styles_to_table(
    table_shape,
    table_data: Dict[str, Any],
    font_name: str = "Pretendard",
    is_dark: bool = False,
    slide_bg_rgb: tuple[int, int, int] = (15, 23, 42),
) -> None:
  tbl = table_shape.table
  _disable_default_pptx_table_style(tbl)
  blend_rgb = (slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2])
  default_border_rgb = RGBColor(51, 65, 85) if is_dark else RGBColor(233, 224, 238)

  if table_data.get("headers") and len(tbl.rows) > 0:
    for c_idx, h in enumerate(table_data["headers"]):
      if c_idx >= len(tbl.rows[0].cells):
        break
      c = tbl.rows[0].cells[c_idx]
      c.fill.solid()
      h_bg = h.get("bgColor") if isinstance(h, dict) else None
      h_col = h.get("color") if isinstance(h, dict) else None
      h_size = float(h.get("fontSizePt") or 11.0) if isinstance(h, dict) else 11.0
      if h_bg:
        h_bg_rgb = parse_color_value(h_bg, default=RGBColor(15, 23, 42) if is_dark else RGBColor(244, 237, 246), bg_blend_rgb=blend_rgb)
      else:
        h_bg_rgb = RGBColor(15, 23, 42) if is_dark else RGBColor(244, 237, 246)
      c.fill.fore_color.rgb = h_bg_rgb
      _apply_cell_bottom_border(c, default_border_rgb, width_emu=12700)

      h_fg_rgb = parse_color_value(h_col, default=RGBColor(148, 163, 184) if is_dark else RGBColor(95, 0, 128))
      for p in c.text_frame.paragraphs:
        _style_paragraph_and_runs(p, font_name, h_size, True, h_fg_rgb)

  for r_idx, row_data in enumerate(table_data.get("rows", [])):
    row_num = r_idx + (1 if table_data.get("headers") else 0)
    if row_num >= len(tbl.rows):
      break
    for c_idx, cell_data in enumerate(row_data):
      if c_idx >= len(tbl.columns):
        break
      cell = tbl.cell(row_num, c_idx)
      cell.fill.solid()

      cell_bg = cell_data.get("bgColor")
      if cell_bg:
        row_bg_rgb = parse_color_value(cell_bg, default=RGBColor(30, 41, 59) if is_dark else RGBColor(255, 255, 255), bg_blend_rgb=blend_rgb)
      elif is_dark:
        row_bg_rgb = RGBColor(30, 41, 59) if r_idx % 2 == 0 else RGBColor(22, 30, 46)
      else:
        row_bg_rgb = RGBColor(255, 255, 255)
      cell.fill.fore_color.rgb = row_bg_rgb
      _apply_cell_bottom_border(cell, default_border_rgb, width_emu=9525)

      default_cell_fg = RGBColor(241, 245, 249) if is_dark else RGBColor(17, 17, 21)
      if "color" in cell_data and cell_data["color"]:
        col_rgb = parse_color_value(cell_data["color"], default=default_cell_fg)
      elif "color_rgb" in cell_data and cell_data["color_rgb"]:
        col_rgb = RGBColor(*cell_data["color_rgb"])
      else:
        col_rgb = default_cell_fg

      bg_lum = row_bg_rgb[0] * 0.299 + row_bg_rgb[1] * 0.587 + row_bg_rgb[2] * 0.114
      fg_lum = col_rgb[0] * 0.299 + col_rgb[1] * 0.587 + col_rgb[2] * 0.114
      if bg_lum < 120 and fg_lum < 90:
        col_rgb = RGBColor(241, 245, 249)
      elif bg_lum >= 150 and fg_lum > 180:
        col_rgb = RGBColor(17, 17, 21)

      if not cell.text.strip() and cell_data.get("text"):
        cell.text = cell_data["text"]

      c_size = float(cell_data.get("fontSizePt") or 10.5)
      for p in cell.text_frame.paragraphs:
        if cell_data.get("has_badge"):
          b_col = parse_color_value(
              cell_data.get("badge_color") or cell_data.get("color"),
              default=RGBColor(96, 165, 250) if is_dark else RGBColor(95, 0, 128),
          )
          _style_paragraph_and_runs(p, font_name, c_size, True, b_col)
        else:
          _style_paragraph_and_runs(p, font_name, c_size, bool(cell_data.get("is_bold")), col_rgb)


def ensure_slide_canvas_background(
    slide,
    prs: Optional[Presentation] = None,
    bg_color: Optional[RGBColor] = None,
    is_dark: bool = False,
) -> None:
  if bg_color is None:
    bg_color = RGBColor(15, 23, 42) if is_dark else RGBColor(255, 255, 255)

  bg_shapes = [
      s for s in slide.shapes
      if s.shape_type == 1 and s.left.inches <= 0.1 and s.top.inches <= 0.1 and s.width.inches >= 13.0 and s.height.inches >= 7.0
  ]
  if len(bg_shapes) > 1:
    for dup in bg_shapes[1:]:
      try:
        slide.shapes._spTree.remove(dup._element)
      except Exception:
        pass

  if bg_shapes:
    full_bleed_bg = bg_shapes[0]
    full_bleed_bg.fill.solid()
    full_bleed_bg.fill.fore_color.rgb = bg_color
    full_bleed_bg.line.fill.background()
  else:
    w = prs.slide_width if prs else Inches(13.333333)
    h = prs.slide_height if prs else Inches(7.5)
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), w, h)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = bg_color
    bg_shape.line.fill.background()
    try:
      spTree = slide.shapes._spTree
      spTree.insert(2, bg_shape._element)
    except Exception:
      pass


def embed_extracted_slide_images(
    slide,
    geom: Optional[Dict[str, Any]],
    screenshot_source: Optional[Union[bytes, str]],
) -> None:
  if not geom or not screenshot_source or not geom.get("images"):
    return

  try:
    if isinstance(screenshot_source, bytes):
      img_stream = io.BytesIO(screenshot_source)
    elif isinstance(screenshot_source, str) and os.path.exists(screenshot_source):
      img_stream = screenshot_source
    else:
      return

    with Image.open(img_stream) as full_img:
      img_w, img_h = full_img.size
      for img_info in geom.get("images", []):
        r = img_info.get("rect", {})
        left_in = float(r.get("left", 0.0))
        top_in = float(r.get("top", 0.0))
        width_in = float(r.get("width", 0.0))
        height_in = float(r.get("height", 0.0))
        if width_in < 0.25 or height_in < 0.25:
          continue

        already_embedded = False
        for s in slide.shapes:
          if getattr(s, "shape_type", None) == 13:
            s_l = s.left.inches if hasattr(s.left, "inches") else s.left / 914400.0
            s_t = s.top.inches if hasattr(s.top, "inches") else s.top / 914400.0
            if abs(s_l - left_in) < 0.35 and abs(s_t - top_in) < 0.35:
              already_embedded = True
              break
        if already_embedded:
          continue

        px_l = max(0, int(round((left_in / 13.333333) * img_w)))
        px_t = max(0, int(round((top_in / 7.5) * img_h)))
        px_r = min(img_w, int(round(((left_in + width_in) / 13.333333) * img_w)))
        px_b = min(img_h, int(round(((top_in + height_in) / 7.5) * img_h)))
        if px_r - px_l < 10 or px_b - px_t < 10:
          continue

        cropped = full_img.crop((px_l, px_t, px_r, px_b))
        bio = io.BytesIO()
        cropped.save(bio, format="PNG")
        bio.seek(0)
        slide.shapes.add_picture(
            bio,
            Inches(left_in),
            Inches(top_in),
            Inches(width_in),
            Inches(height_in),
        )
  except Exception as e:
    logger.warning("Failed to embed extracted slide images: %s", e)


def build_slide_from_geometry(
    slide,
    geom: Dict[str, Any],
    font_name: str = "Pretendard",
    brand_color_rgb: Optional[RGBColor] = None,
    screenshot_source: Optional[Union[bytes, str]] = None,
) -> None:
  if not geom:
    return

  if brand_color_rgb is None:
    brand_color_rgb = RGBColor(37, 99, 235)

  bg_color = parse_color_value(geom.get("slideBgColor"), RGBColor(255, 255, 255))
  is_dark = (bg_color[0] * 0.299 + bg_color[1] * 0.587 + bg_color[2] * 0.114) < 128
  ensure_slide_canvas_background(slide, prs=None, bg_color=bg_color, is_dark=is_dark)
  blend_rgb = (bg_color[0], bg_color[1], bg_color[2])

  if geom.get("headerDividerTop"):
    h_div = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.35),
        Inches(float(geom["headerDividerTop"])),
        Inches(12.63),
        Inches(0.015),
    )
    h_div.fill.solid()
    h_div.fill.fore_color.rgb = parse_color_value(
        geom.get("headerDividerColor"),
        RGBColor(30, 41, 59) if is_dark else RGBColor(235, 226, 240),
        bg_blend_rgb=blend_rgb,
    )
    h_div.line.fill.background()

  for b_data in geom.get("headerBadges", []):
    r = b_data["rect"]
    b_txt = b_data["text"]
    raw_bg = (b_data.get("styles", {}).get("backgroundColor") or "").strip()
    has_bg = bool(raw_bg and raw_bg not in ("transparent", "rgba(0, 0, 0, 0)", "none"))
    bg_rgb = parse_color_value(raw_bg if has_bg else None, None, bg_blend_rgb=blend_rgb)
    fg_rgb = parse_color_value(b_data["styles"]["color"], RGBColor(255, 255, 255) if has_bg else brand_color_rgb)
    b_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(r["left"]),
        Inches(r["top"]),
        Inches(max(0.8, r["width"])),
        Inches(max(0.22, r["height"])),
    )
    if bg_rgb is not None:
      b_shape.fill.solid()
      b_shape.fill.fore_color.rgb = bg_rgb
    else:
      b_shape.fill.background()
    b_border = parse_color_value(b_data["styles"].get("borderColor"), None, bg_blend_rgb=blend_rgb)
    if b_border:
      b_shape.line.color.rgb = b_border
      b_shape.line.width = Pt(1.0)
    else:
      b_shape.line.fill.background()
    try:
      b_shape.adjustments[0] = 0.5
    except Exception:
      pass
    tf = b_shape.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.08 if has_bg else 0.0)
    tf.margin_right = Inches(0.08 if has_bg else 0.0)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = b_txt
    p.font.name = font_name
    p.font.size = Pt(b_data["styles"].get("fontSizePt", 10.0))
    p.font.bold = True
    p.font.color.rgb = fg_rgb
    p.alignment = PP_ALIGN.CENTER if has_bg else PP_ALIGN.LEFT

  if geom.get("num") and geom["num"]["text"]:
    n_data = geom["num"]
    r = n_data["rect"]
    num_box = slide.shapes.add_textbox(
        Inches(min(12.2, r["left"])),
        Inches(r["top"]),
        Inches(max(0.8, r["width"])),
        Inches(max(0.25, r["height"])),
    )
    p = num_box.text_frame.paragraphs[0]
    p.text = n_data["text"]
    p.font.name = font_name
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = parse_color_value(n_data["styles"].get("color") if n_data.get("styles") else None, RGBColor(148, 163, 184) if is_dark else RGBColor(100, 116, 139))
    p.alignment = PP_ALIGN.RIGHT

  if geom.get("title") and geom["title"]["text"]:
    t_data = geom["title"]
    r = t_data["rect"]
    st = t_data["styles"]
    is_center = st.get("textAlign") == "center"
    t_left = Inches(r["left"]) if is_center else Inches(max(0.35, r["left"]))
    t_top = Inches(r["top"])
    t_w = Inches(r["width"]) if is_center else Inches(min(12.4, max(r["width"], 8.0)))
    t_h = Inches(max(0.42, r["height"]))
    t_box = slide.shapes.add_textbox(t_left, t_top, t_w, t_h)
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    t_runs = t_data.get("runs") or []
    default_t_col = parse_color_value(st.get("color"), RGBColor(248, 250, 252) if is_dark else brand_color_rgb)
    if len(t_runs) > 1:
      p.text = ""
      for r_info in t_runs:
        run = p.add_run()
        run.text = r_info.get("text", "")
        run.font.name = font_name
        run.font.size = Pt(st.get("fontSizePt", 22.0))
        run.font.bold = True
        run.font.color.rgb = parse_color_value(r_info.get("color"), default_t_col)
    else:
      p.text = t_data["text"]
      p.font.name = font_name
      p.font.size = Pt(st.get("fontSizePt", 22.0))
      p.font.bold = True
      p.font.color.rgb = default_t_col
    if is_center:
      p.alignment = PP_ALIGN.CENTER

  if geom.get("sub") and geom["sub"]["text"]:
    s_data = geom["sub"]
    r = s_data["rect"]
    st = s_data["styles"]
    is_center = st.get("textAlign") == "center"
    s_left = Inches(r["left"]) if is_center else Inches(max(0.35, r["left"]))
    s_top = Inches(r["top"])
    s_w = Inches(r["width"]) if is_center else Inches(min(12.4, max(r["width"], 8.0)))
    s_h = Inches(max(0.30, r["height"]))
    s_box = slide.shapes.add_textbox(s_left, s_top, s_w, s_h)
    tf = s_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = s_data["text"]
    p.font.name = font_name
    p.font.size = Pt(st.get("fontSizePt", 12.5))
    p.font.bold = bool(st.get("isBold"))
    p.font.color.rgb = parse_color_value(st.get("color"), RGBColor(148, 163, 184) if is_dark else RGBColor(100, 116, 139))
    if is_center:
      p.alignment = PP_ALIGN.CENTER

  for d in geom.get("desc", []):
    r = d["rect"]
    st = d["styles"]
    d_box = slide.shapes.add_textbox(Inches(r["left"]), Inches(r["top"]), Inches(max(2.0, r["width"])), Inches(max(0.32, r["height"])))
    tf = d_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = d["text"]
    p.font.name = font_name
    p.font.size = Pt(st.get("fontSizePt", 12.0))
    p.font.color.rgb = parse_color_value(st.get("color"), RGBColor(203, 213, 225) if is_dark else RGBColor(74, 77, 82))

  for card in geom.get("cards", []):
    r = card["rect"]
    c_left = Inches(r["left"])
    c_top = Inches(r["top"])
    c_w = Inches(r["width"])
    c_h = Inches(r["height"])

    if card.get("isBridge") and card.get("bridge"):
      b_info = card["bridge"]
      br = b_info["rect"]
      b_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(br["left"]), Inches(br["top"]), Inches(br["width"]), Inches(br["height"]))
      b_shape.fill.solid()
      b_shape.fill.fore_color.rgb = parse_color_value(b_info["styles"]["backgroundColor"], brand_color_rgb)
      b_shape.line.fill.background()
      tf = b_shape.text_frame
      tf.word_wrap = True
      tf.margin_left = Inches(0.05)
      tf.margin_right = Inches(0.05)
      p0 = tf.paragraphs[0]
      p0.text = b_info.get("arrow", "➔")
      p0.font.name = font_name
      p0.font.size = Pt(15.0)
      p0.font.bold = True
      p0.font.color.rgb = RGBColor(255, 255, 255)
      p0.alignment = PP_ALIGN.CENTER
      p1 = tf.add_paragraph()
      p1.text = b_info.get("text", "")
      p1.font.name = font_name
      p1.font.size = Pt(9.0)
      p1.font.bold = True
      p1.font.color.rgb = RGBColor(255, 255, 255)
      p1.alignment = PP_ALIGN.CENTER
      continue

    raw_c_bg = (card.get("styles", {}).get("backgroundColor") or "").strip()
    raw_c_border = (card.get("styles", {}).get("borderColor") or "").strip()
    has_c_bg = bool(raw_c_bg and raw_c_bg not in ("transparent", "rgba(0, 0, 0, 0)", "none"))
    has_c_border = bool(raw_c_border and raw_c_border not in ("transparent", "rgba(0, 0, 0, 0)", "none"))

    c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, c_w, c_h)
    if has_c_bg:
      c_shape.fill.solid()
      c_bg = parse_color_value(raw_c_bg, RGBColor(30, 41, 59) if is_dark else RGBColor(255, 255, 255), bg_blend_rgb=blend_rgb)
      c_shape.fill.fore_color.rgb = c_bg
    else:
      c_shape.fill.background()

    if has_c_border:
      c_border = parse_color_value(raw_c_border, RGBColor(51, 65, 85) if is_dark else RGBColor(226, 232, 240), bg_blend_rgb=blend_rgb)
      c_shape.line.color.rgb = c_border
      c_shape.line.width = Pt(1.0)
    else:
      c_shape.line.fill.background()

    if card.get("isAttachCard") and card.get("attachData"):
      ad = card["attachData"]
      tf = c_shape.text_frame
      tf.word_wrap = True
      tf.margin_left = Inches(0.18)
      tf.margin_right = Inches(0.18)
      tf.margin_top = Inches(0.14)
      tf.margin_bottom = Inches(0.14)
      p0 = tf.paragraphs[0]
      p0.text = ad["target"] + (f"  [{ad['mode']}]" if ad.get("mode") else "")
      p0.font.name = font_name
      p0.font.size = Pt(12.5)
      p0.font.bold = True
      p0.font.color.rgb = brand_color_rgb
      if ad.get("mech"):
        p1 = tf.add_paragraph()
        p1.text = ad["mech"]
        p1.font.name = "Consolas"
        p1.font.size = Pt(9.0)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(24, 107, 242)
      if ad.get("desc"):
        p2 = tf.add_paragraph()
        p2.text = ad["desc"]
        p2.font.name = font_name
        p2.font.size = Pt(10.0)
        p2.font.color.rgb = RGBColor(74, 85, 104)
      continue

    curr_top_in = r["top"] + (0.18 if card.get("hasTopAccent") else 0.12)
    if card.get("hasTopAccent"):
      bar_color = parse_color_value(card.get("topAccentColor"), brand_color_rgb)
      accent_bar = slide.shapes.add_shape(
          MSO_SHAPE.ROUNDED_RECTANGLE,
          c_left + Inches(0.20),
          c_top + Inches(0.08),
          c_w - Inches(0.40),
          Inches(0.06),
      )
      accent_bar.fill.solid()
      accent_bar.fill.fore_color.rgb = bar_color
      accent_bar.line.fill.background()
      try:
        accent_bar.adjustments[0] = 0.5
      except Exception:
        pass

    if card.get("hasLeftAccent"):
      l_bar_color = parse_color_value(card.get("leftAccentColor"), brand_color_rgb)
      l_bar = slide.shapes.add_shape(
          MSO_SHAPE.ROUNDED_RECTANGLE,
          c_left + Inches(0.06),
          c_top + Inches(0.12),
          Inches(0.06),
          max(Inches(0.24), c_h - Inches(0.24)),
      )
      l_bar.fill.solid()
      l_bar.fill.fore_color.rgb = l_bar_color
      l_bar.line.fill.background()

    if card.get("tag") and card["tag"]["text"]:
      tg = card["tag"]
      tg_r = tg["rect"]
      tg_raw_bg = (tg.get("styles", {}).get("backgroundColor") or "").strip()
      tg_has_bg = bool(tg_raw_bg and tg_raw_bg not in ("transparent", "rgba(0, 0, 0, 0)", "none"))
      tg_bg = parse_color_value(tg_raw_bg if tg_has_bg else None, None, bg_blend_rgb=blend_rgb)
      tg_fg = parse_color_value(tg["styles"]["color"], RGBColor(255, 255, 255) if tg_has_bg else brand_color_rgb)
      tg_top_in = max(curr_top_in, tg_r.get("top", curr_top_in))
      tg_h_in = max(0.22, tg_r.get("height", 0.22))
      tg_shape = slide.shapes.add_shape(
          MSO_SHAPE.ROUNDED_RECTANGLE,
          Inches(tg_r.get("left", r["left"] + 0.18)),
          Inches(tg_top_in),
          Inches(max(0.8, tg_r["width"])),
          Inches(tg_h_in),
      )
      if tg_bg is not None:
        tg_shape.fill.solid()
        tg_shape.fill.fore_color.rgb = tg_bg
      else:
        tg_shape.fill.background()
      tg_shape.line.fill.background()
      tf = tg_shape.text_frame
      tf.word_wrap = False
      tf.margin_left = Inches(0.06 if tg_has_bg else 0.0)
      tf.margin_right = Inches(0.06 if tg_has_bg else 0.0)
      p = tf.paragraphs[0]
      p.text = tg["text"]
      p.font.name = font_name
      p.font.size = Pt(tg["styles"].get("fontSizePt", 9.5))
      p.font.bold = True
      p.font.color.rgb = tg_fg
      curr_top_in = tg_top_in + tg_h_in + 0.04

    for b in card.get("badges", []):
      br = b.get("rect", {})
      b_txt = b.get("text", "")
      if b_txt and br:
        b_w_in = max(0.60, br.get("width", 1.0))
        b_w = Inches(b_w_in)
        b_h = Inches(max(0.22, br.get("height", 0.24)))
        b_left = Inches(br.get("left", r.get("left", 0.0) + r.get("width", 1.0) - b_w_in - 0.15))
        b_top = Inches(br.get("top", r.get("top", 0.0) + 0.18))
        b_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            b_left,
            b_top,
            b_w,
            b_h,
        )
        b_shape.fill.solid()
        b_bg = parse_color_value(b.get("styles", {}).get("backgroundColor"), RGBColor(241, 245, 249), bg_blend_rgb=blend_rgb)
        b_shape.fill.fore_color.rgb = b_bg
        b_border = parse_color_value(b.get("styles", {}).get("borderColor"), RGBColor(226, 232, 240))
        b_shape.line.color.rgb = b_border
        b_shape.line.width = Pt(1.0)
        try:
          b_shape.adjustments[0] = 0.35
        except Exception:
          pass
        tf = b_shape.text_frame
        tf.word_wrap = False
        tf.margin_left = Inches(0.06)
        tf.margin_right = Inches(0.06)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        p = tf.paragraphs[0]
        p.text = b_txt
        p.font.name = font_name
        p.font.size = Pt(b.get("styles", {}).get("fontSizePt", 9.0))
        p.font.bold = True
        p.font.color.rgb = parse_color_value(b.get("styles", {}).get("color"), RGBColor(78, 86, 95))
        p.alignment = PP_ALIGN.CENTER

    if card.get("title") and card["title"] != "➔":
      t_rect = card.get("titleRect") or {}
      t_left_in = t_rect.get("left", r["left"] + 0.18)
      t_top_in = max(curr_top_in, t_rect.get("top", curr_top_in))
      t_h_in = max(0.32, t_rect.get("height", 0.34))
      same_row_right_lefts = [
          b["rect"]["left"]
          for b in card.get("badges", [])
          if b.get("rect")
          and abs(b["rect"].get("top", 0.0) - t_top_in) < 0.28
          and b["rect"].get("left", 0.0) > t_left_in + 0.5
      ] + [
          sc["rect"]["left"]
          for sc in card.get("subCards", [])
          if sc.get("rect")
          and abs(sc["rect"].get("top", 0.0) - t_top_in) < 0.45
          and sc["rect"].get("left", 0.0) > t_left_in + 0.5
      ]
      b_first_left = min(same_row_right_lefts, default=None)
      if b_first_left is not None:
        title_w_in = max(1.0, b_first_left - t_left_in - 0.08)
      else:
        title_w_in = max(1.2, (r["left"] + r["width"] - 0.16) - t_left_in)
      t_box = slide.shapes.add_textbox(
          Inches(t_left_in),
          Inches(t_top_in),
          Inches(title_w_in),
          Inches(t_h_in),
      )
      tf = t_box.text_frame
      tf.word_wrap = True
      tf.margin_left = Inches(0.02)
      tf.margin_right = Inches(0.02)
      tf.margin_top = Inches(0.02)
      tf.margin_bottom = Inches(0.02)
      p = tf.paragraphs[0]
      p.text = card["title"]
      p.font.name = font_name
      p.font.size = Pt(card.get("titleStyles", {}).get("fontSizePt", 13.0) if card.get("titleStyles") else 13.0)
      p.font.bold = True
      p.font.color.rgb = parse_color_value(card.get("titleStyles", {}).get("color") if card.get("titleStyles") else None, RGBColor(96, 165, 250) if is_dark else brand_color_rgb)
      curr_top_in = t_top_in + t_h_in + 0.04

    if card.get("desc") and card["desc"]["text"]:
      d_rect = card["desc"].get("rect") or {}
      d_left_in = d_rect.get("left", r["left"] + 0.18)
      d_top_in = max(curr_top_in, d_rect.get("top", curr_top_in))
      d_h_in = max(0.28, d_rect.get("height", 0.32))
      d_w_in = max(1.0, (r["left"] + r["width"] - 0.16) - d_left_in)
      d_box = slide.shapes.add_textbox(Inches(d_left_in), Inches(d_top_in), Inches(d_w_in), Inches(d_h_in))
      tf = d_box.text_frame
      tf.word_wrap = True
      tf.margin_left = Inches(0.02)
      tf.margin_right = Inches(0.02)
      tf.margin_top = Inches(0.01)
      tf.margin_bottom = Inches(0.01)
      p = tf.paragraphs[0]
      p.text = card["desc"]["text"]
      p.font.name = font_name
      p.font.size = Pt(card["desc"]["styles"].get("fontSizePt", 10.5))
      p.font.color.rgb = parse_color_value(card["desc"]["styles"].get("color"), RGBColor(148, 163, 184) if is_dark else RGBColor(100, 116, 139))
      curr_top_in = d_top_in + d_h_in + 0.04

    for ps in card.get("pipelineSteps", []):
      psr = ps["rect"]
      ps_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(psr["left"]), Inches(psr["top"]), Inches(psr["width"]), Inches(psr["height"]))
      ps_shape.fill.solid()
      ps_bg = parse_color_value(ps["styles"]["backgroundColor"], RGBColor(255, 255, 255), bg_blend_rgb=blend_rgb)
      ps_shape.fill.fore_color.rgb = ps_bg
      ps_border = parse_color_value(ps["styles"].get("borderColor"), RGBColor(217, 226, 236), bg_blend_rgb=blend_rgb)
      ps_shape.line.color.rgb = ps_border
      ps_shape.line.width = Pt(1.0)

      num_r = ps.get("numRect")
      num_st = ps.get("numStyles") or {}
      t_st = ps.get("titleStyles") or ps.get("styles") or {}
      d_st = ps.get("descStyles") or {}
      if num_r and ps.get("stepNum"):
        n_dim = max(0.28, min(num_r.get("width", 0.32), num_r.get("height", 0.32)))
        n_shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(num_r.get("left", psr["left"] + 0.12)),
            Inches(num_r.get("top", psr["top"] + (psr["height"] - n_dim) / 2.0)),
            Inches(n_dim),
            Inches(n_dim),
        )
        n_shape.fill.solid()
        n_shape.fill.fore_color.rgb = parse_color_value(num_st.get("backgroundColor"), brand_color_rgb, bg_blend_rgb=blend_rgb)
        n_shape.line.fill.background()
        n_tf = n_shape.text_frame
        n_tf.word_wrap = False
        n_tf.margin_left = Inches(0.0)
        n_tf.margin_right = Inches(0.0)
        n_tf.margin_top = Inches(0.0)
        n_tf.margin_bottom = Inches(0.0)
        n_p = n_tf.paragraphs[0]
        n_p.text = str(ps["stepNum"])
        n_p.font.name = font_name
        n_p.font.size = Pt(num_st.get("fontSizePt", 10.0))
        n_p.font.bold = True
        n_p.font.color.rgb = RGBColor(255, 255, 255)
        n_p.alignment = PP_ALIGN.CENTER

        tx_left_in = (ps.get("titleRect") or {}).get("left", num_r.get("left", psr["left"] + 0.12) + n_dim + 0.12)
        tx_w_in = max(1.0, (psr["left"] + psr["width"] - 0.12) - tx_left_in)
        tx_box = slide.shapes.add_textbox(
            Inches(tx_left_in),
            Inches(psr["top"] + 0.06),
            Inches(tx_w_in),
            Inches(max(0.32, psr["height"] - 0.12)),
        )
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.02)
        tf.margin_right = Inches(0.02)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
      else:
        tf = ps_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.10)
        tf.margin_right = Inches(0.10)

      p0 = tf.paragraphs[0]
      p0.text = ps["title"]
      p0.font.name = font_name
      p0.font.size = Pt(t_st.get("fontSizePt", 11.0))
      p0.font.bold = True
      p0.font.color.rgb = parse_color_value(t_st.get("color"), brand_color_rgb)
      if ps.get("desc"):
        p1 = tf.add_paragraph()
        p1.text = ps["desc"]
        p1.font.name = font_name
        p1.font.size = Pt(d_st.get("fontSizePt", 9.5))
        p1.font.color.rgb = parse_color_value(d_st.get("color"), RGBColor(51, 65, 85))

    for fl in card.get("flows", []):
      for node in fl.get("nodes", []):
        nr = node["rect"]
        if node.get("isArrow"):
          arr_box = slide.shapes.add_textbox(Inches(nr["left"]), Inches(nr["top"]), Inches(nr["width"]), Inches(nr["height"]))
          tf = arr_box.text_frame
          tf.word_wrap = False
          p = tf.paragraphs[0]
          p.text = node["text"]
          p.font.name = font_name
          p.font.size = Pt(12.0)
          p.font.bold = True
          p.font.color.rgb = brand_color_rgb
          p.alignment = PP_ALIGN.CENTER
        else:
          n_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(nr["left"]), Inches(nr["top"]), Inches(nr["width"]), Inches(nr["height"]))
          n_shape.fill.solid()
          n_bg = parse_color_value(node["styles"]["backgroundColor"], RGBColor(241, 245, 249) if not node.get("isHighlight") else RGBColor(238, 242, 255), bg_blend_rgb=blend_rgb)
          n_shape.fill.fore_color.rgb = n_bg
          n_border = parse_color_value(node["styles"].get("borderColor"), brand_color_rgb if node.get("isHighlight") else RGBColor(203, 213, 225))
          n_shape.line.color.rgb = n_border
          n_shape.line.width = Pt(1.0)
          tf = n_shape.text_frame
          tf.word_wrap = True
          tf.margin_left = Inches(0.06)
          tf.margin_right = Inches(0.06)
          p = tf.paragraphs[0]
          p.text = node["text"]
          p.font.name = font_name
          p.font.size = Pt(node["styles"].get("fontSizePt", 9.5))
          p.font.bold = True
          p.font.color.rgb = parse_color_value(node["styles"].get("color"), brand_color_rgb if node.get("isHighlight") else RGBColor(30, 41, 59))
          p.alignment = PP_ALIGN.CENTER

    for fi in card.get("fItems", []):
      fir = fi["rect"]
      fi_st = fi.get("styles") or {}
      fi_raw_bg = (fi_st.get("backgroundColor") or "").strip()
      fi_has_bg = bool(fi_raw_bg and fi_raw_bg not in ("transparent", "rgba(0, 0, 0, 0)", "none") and fi_raw_bg != raw_c_bg)
      ic_r = fi.get("iconRect")
      tx_r = fi.get("textRect")
      has_icon = bool(fi.get("icon") and ic_r)

      if fi_has_bg and not has_icon:
        fi_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(fir["left"]),
            Inches(fir["top"]),
            Inches(fir["width"]),
            Inches(max(0.32, fir["height"])),
        )
        fi_shape.fill.solid()
        fi_bg = parse_color_value(
            fi_raw_bg,
            RGBColor(15, 23, 42) if is_dark else RGBColor(248, 250, 252),
            bg_blend_rgb=blend_rgb,
        )
        fi_shape.fill.fore_color.rgb = fi_bg
        fi_border = parse_color_value(
            fi_st.get("borderColor"),
            RGBColor(51, 65, 85) if is_dark else RGBColor(226, 232, 240),
            bg_blend_rgb=blend_rgb,
        )
        fi_shape.line.color.rgb = fi_border
        fi_shape.line.width = Pt(1.0)
        tf = fi_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.06)
        tf.margin_bottom = Inches(0.06)
        if fi.get("title") and fi.get("desc"):
          t_st = fi.get("titleStyles") or fi.get("textStyles") or {}
          d_st = fi.get("descStyles") or fi.get("textStyles") or {}
          p0 = tf.paragraphs[0]
          p0.text = fi["title"]
          p0.font.name = font_name
          p0.font.size = Pt(t_st.get("fontSizePt", 10.5))
          p0.font.bold = True
          p0.font.color.rgb = parse_color_value(t_st.get("color"), RGBColor(248, 250, 252) if is_dark else RGBColor(15, 23, 42))
          p1 = tf.add_paragraph()
          p1.text = fi["desc"]
          p1.font.name = font_name
          p1.font.size = Pt(d_st.get("fontSizePt", 9.5))
          p1.font.color.rgb = parse_color_value(d_st.get("color"), RGBColor(203, 213, 225) if is_dark else RGBColor(71, 85, 105))
        elif fi.get("text"):
          p0 = tf.paragraphs[0]
          p0.text = fi["text"]
          p0.font.name = font_name
          p0.font.size = Pt(fi.get("textStyles", {}).get("fontSizePt", 10.0) if fi.get("textStyles") else 10.0)
          p0.font.color.rgb = parse_color_value(
              fi.get("textStyles", {}).get("color") if fi.get("textStyles") else None,
              RGBColor(17, 17, 21),
          )
        continue

      if fi_has_bg:
        fi_bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(fir["left"]),
            Inches(fir["top"]),
            Inches(fir["width"]),
            Inches(max(0.32, fir["height"])),
        )
        fi_bg_shape.fill.solid()
        fi_bg_shape.fill.fore_color.rgb = parse_color_value(
            fi_raw_bg,
            RGBColor(15, 23, 42) if is_dark else RGBColor(248, 250, 252),
            bg_blend_rgb=blend_rgb,
        )
        fi_bg_shape.line.color.rgb = parse_color_value(
            fi_st.get("borderColor"),
            RGBColor(51, 65, 85) if is_dark else RGBColor(226, 232, 240),
            bg_blend_rgb=blend_rgb,
        )
        fi_bg_shape.line.width = Pt(1.0)

      ic_w = 0.0
      ic_left = 0.0
      if has_icon:
        ic_w = max(0.38, ic_r.get("width", 0.40))
        ic_h = max(0.22, ic_r.get("height", 0.24))
        ic_left = ic_r.get("left", fir.get("left", r.get("left", 0.0) + 0.18))
        ic_top = ic_r.get("top", fir.get("top", r.get("top", 0.0) + 0.18))
        ic_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(ic_left),
            Inches(ic_top),
            Inches(ic_w),
            Inches(ic_h),
        )
        ic_shape.fill.solid()
        ic_bg = parse_color_value(
            fi.get("iconStyles", {}).get("backgroundColor") if fi.get("iconStyles") else None,
            brand_color_rgb,
            bg_blend_rgb=blend_rgb,
        )
        ic_shape.fill.fore_color.rgb = ic_bg
        ic_shape.line.fill.background()
        tf = ic_shape.text_frame
        tf.word_wrap = False
        tf.margin_left = Inches(0.04)
        tf.margin_right = Inches(0.04)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        p = tf.paragraphs[0]
        p.text = fi["icon"]
        p.font.name = font_name
        p.font.size = Pt(fi.get("iconStyles", {}).get("fontSizePt", 9.0) if fi.get("iconStyles") else 9.0)
        p.font.bold = True
        p.font.color.rgb = parse_color_value(
            fi.get("iconStyles", {}).get("color") if fi.get("iconStyles") else None,
            RGBColor(255, 255, 255),
        )
        p.alignment = PP_ALIGN.CENTER

      if fi.get("text"):
        card_r = r.get("left", 0.0) + r.get("width", 4.0) - 0.15
        if ic_w > 0:
          min_tx_left = ic_left + ic_w + 0.08
          tx_left = max(tx_r.get("left", min_tx_left) if tx_r else min_tx_left, min_tx_left)
        else:
          tx_left = tx_r.get("left", fir.get("left", r.get("left", 0.0) + 0.18)) if tx_r else fir.get("left", r.get("left", 0.0) + 0.18)

        tx_top = tx_r.get("top", fir.get("top", r.get("top", 0.0) + 0.18)) if tx_r else fir.get("top", r.get("top", 0.0) + 0.18)
        tx_w = max(0.8, card_r - tx_left)
        tx_h = max(0.24, fir.get("height", 0.24) if (fi.get("title") and fi.get("desc")) else (tx_r.get("height", 0.24) if tx_r else 0.24))

        tx_box = slide.shapes.add_textbox(Inches(tx_left), Inches(tx_top), Inches(tx_w), Inches(tx_h))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.02)
        tf.margin_right = Inches(0.02)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        if fi.get("title") and fi.get("desc"):
          t_st = fi.get("titleStyles") or fi.get("textStyles") or {}
          d_st = fi.get("descStyles") or fi.get("textStyles") or {}
          p0 = tf.paragraphs[0]
          p0.text = fi["title"]
          p0.font.name = font_name
          p0.font.size = Pt(t_st.get("fontSizePt", 10.5))
          p0.font.bold = True
          p0.font.color.rgb = parse_color_value(t_st.get("color"), RGBColor(248, 250, 252) if is_dark else RGBColor(15, 23, 42))
          p1 = tf.add_paragraph()
          p1.text = fi["desc"]
          p1.font.name = font_name
          p1.font.size = Pt(d_st.get("fontSizePt", 9.5))
          p1.font.color.rgb = parse_color_value(d_st.get("color"), RGBColor(203, 213, 225) if is_dark else RGBColor(71, 85, 105))
        else:
          p = tf.paragraphs[0]
          p.text = fi["text"]
          p.font.name = font_name
          p.font.size = Pt(fi.get("textStyles", {}).get("fontSizePt", 10.0) if fi.get("textStyles") else 10.0)
          p.font.color.rgb = parse_color_value(
              fi.get("textStyles", {}).get("color") if fi.get("textStyles") else None,
              RGBColor(17, 17, 21),
          )

    for sc in card.get("subCards", []):
      scr = sc["rect"]
      sc_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(scr["left"]), Inches(scr["top"]), Inches(scr["width"]), Inches(scr["height"]))
      sc_shape.fill.solid()
      sc_bg = parse_color_value(sc["styles"]["backgroundColor"], RGBColor(15, 23, 42) if is_dark else RGBColor(248, 250, 252), bg_blend_rgb=blend_rgb)
      sc_shape.fill.fore_color.rgb = sc_bg
      if sc.get("hasLeftAccent"):
        sc_shape.line.fill.background()
        sc_l_bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(scr["left"] + 0.04),
            Inches(scr["top"] + 0.06),
            Inches(0.05),
            Inches(max(0.16, scr["height"] - 0.12)),
        )
        sc_l_bar.fill.solid()
        sc_l_bar.fill.fore_color.rgb = parse_color_value(sc.get("leftAccentColor"), brand_color_rgb)
        sc_l_bar.line.fill.background()
      else:
        sc_border = parse_color_value(sc["styles"].get("borderColor"), RGBColor(51, 65, 85) if is_dark else RGBColor(226, 232, 240), bg_blend_rgb=blend_rgb)
        sc_shape.line.color.rgb = sc_border
        sc_shape.line.width = Pt(1.0)
      tf = sc_shape.text_frame
      tf.word_wrap = True
      tf.margin_left = Inches(0.16 if sc.get("hasLeftAccent") else 0.12)
      tf.margin_right = Inches(0.12)
      tf.margin_top = Inches(0.06)
      tf.margin_bottom = Inches(0.06)
      p_idx = 0
      if sc.get("head"):
        p0 = tf.paragraphs[0]
        p0.text = sc["head"] + (f"  [{sc['badge']}]" if sc.get("badge") else "")
        p0.font.name = font_name
        head_st = sc.get("headStyles") or {}
        p0.font.size = Pt(head_st.get("fontSizePt", 10.5))
        p0.font.bold = True
        head_col_str = head_st.get("color") or sc["styles"].get("color")
        p0.font.color.rgb = parse_color_value(head_col_str, RGBColor(248, 250, 252) if is_dark else brand_color_rgb)
        p_idx += 1
      sc_items = sc.get("items", [])
      sc_item_styles = sc.get("itemStyles") or []
      for it_idx, it in enumerate(sc_items):
        p_it = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
        use_bullet = len(sc_items) > 1 and not it.startswith(("•", "-", "1", "2", "3", "KPI"))
        p_it.text = f"• {it}" if use_bullet else it
        p_it.font.name = font_name
        it_st = sc_item_styles[it_idx] if it_idx < len(sc_item_styles) else sc["styles"]
        p_it.font.size = Pt(it_st.get("fontSizePt", 9.5))
        p_it.font.color.rgb = parse_color_value(it_st.get("color"), RGBColor(203, 213, 225) if is_dark else RGBColor(30, 41, 59))
        p_idx += 1

    for cb in card.get("codeBlocks", []):
      cbr = cb["rect"]
      cb_h = Inches(cbr["height"]) if cbr.get("height", 0) > 0.3 else Inches(0.40)
      cb_box = slide.shapes.add_shape(
          MSO_SHAPE.ROUNDED_RECTANGLE,
          Inches(cbr["left"]),
          Inches(cbr["top"]),
          Inches(cbr["width"]),
          cb_h,
      )
      cb_box.fill.solid()
      cb_bg = parse_color_value(cb["styles"]["backgroundColor"], RGBColor(15, 23, 42) if is_dark else RGBColor(241, 245, 249), bg_blend_rgb=blend_rgb)
      cb_box.fill.fore_color.rgb = cb_bg
      cb_border = parse_color_value(cb["styles"].get("borderColor"), RGBColor(51, 65, 85) if is_dark else RGBColor(226, 232, 240))
      cb_box.line.color.rgb = cb_border
      cb_box.line.width = Pt(1.0)
      tf = cb_box.text_frame
      tf.word_wrap = True
      tf.margin_left = Inches(0.08)
      tf.margin_right = Inches(0.08)
      tf.margin_top = Inches(0.04)
      tf.margin_bottom = Inches(0.04)
      for p_idx, line in enumerate(cb["text"].split("\n")):
        p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(cb["styles"].get("fontSizePt", 9.0))
        p.font.color.rgb = parse_color_value(cb["styles"].get("color"), RGBColor(203, 213, 225) if is_dark else RGBColor(30, 41, 59))

    paras = card.get("paragraphs", [])
    if paras and not card.get("pipelineSteps"):
      running_p_top = curr_top_in
      card_right_in = r["left"] + r["width"] - 0.14
      for para in paras:
        pr = para.get("rect") or {}
        p_left_in = pr.get("left", r["left"] + 0.18)
        p_top_in = max(running_p_top, pr.get("top", running_p_top))
        max_avail_pw = max(0.8, card_right_in - p_left_in)
        p_w_in = max(0.8, min(max_avail_pw, pr.get("width", max_avail_pw) + 0.12))
        p_h_in = max(0.24, pr.get("height", 0.28))
        p_box = slide.shapes.add_textbox(
            Inches(p_left_in),
            Inches(p_top_in),
            Inches(p_w_in),
            Inches(p_h_in),
        )
        tf = p_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.02)
        tf.margin_right = Inches(0.02)
        tf.margin_top = Inches(0.01)
        tf.margin_bottom = Inches(0.01)
        raw_ptxt = para["text"]
        lines = raw_ptxt.split("\n")
        for l_idx, l_str in enumerate(lines):
          p = tf.paragraphs[0] if l_idx == 0 else tf.add_paragraph()
          use_bullet = l_idx == 0 and bool(para.get("isList")) and not l_str.startswith(("•", "-", "1", "2"))
          p.text = f"• {l_str}" if use_bullet else l_str
          p.font.name = font_name
          p.font.size = Pt(para.get("styles", {}).get("fontSizePt", 10.0))
          p.font.bold = bool(para.get("styles", {}).get("isBold"))
          p.font.color.rgb = parse_color_value(
              para.get("styles", {}).get("color"),
              RGBColor(226, 232, 240) if is_dark else RGBColor(17, 17, 21),
          )
        running_p_top = p_top_in + p_h_in + 0.02

  for tbl in geom.get("tables", []):
    if tbl.get("num_rows", 0) > 0 and tbl.get("num_cols", 0) > 0:
      tr = tbl.get("rect", {})
      tbl_copy = dict(tbl)
      tbl_copy["is_dark"] = is_dark
      tbl_copy["slide_bg_rgb"] = blend_rgb
      tbl_copy["skip_card_wrapper"] = False
      build_styled_native_table(
          slide,
          tbl_copy,
          Inches(tr.get("left", 0.6)),
          Inches(tr.get("top", 1.6)),
          Inches(tr.get("width", 12.1)),
          Inches(tr.get("height", 4.5)),
          font_name=font_name,
      )

  for ch in geom.get("charts", []):
    build_styled_native_chart(
        slide,
        ch,
        font_name=font_name,
        is_dark=is_dark,
        slide_bg_rgb=blend_rgb,
    )

  for hl in geom.get("highlightBoxes", []):
    hl_r = hl.get("rect", {})
    hl_st = hl.get("styles", {})
    hl_txt_st = hl.get("textStyles") or hl_st
    sub_b = hl.get("subBadge")
    hl_left = hl_r.get("left", 0.8)
    hl_top = hl_r.get("top", 6.3)
    hl_w = hl_r.get("width", 11.733)
    hl_h = max(0.36, hl_r.get("height", 0.42))
    hl_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(hl_left),
        Inches(hl_top),
        Inches(hl_w),
        Inches(hl_h),
    )
    hl_shape.fill.solid()
    hl_bg = parse_color_value(
        hl_st.get("backgroundColor"),
        RGBColor(30, 41, 59) if is_dark else RGBColor(241, 245, 249),
        bg_blend_rgb=blend_rgb,
    )
    hl_shape.fill.fore_color.rgb = hl_bg
    if hl.get("hasLeftAccent"):
      hl_shape.line.fill.background()
      hl_l_bar = slide.shapes.add_shape(
          MSO_SHAPE.ROUNDED_RECTANGLE,
          Inches(hl_left + 0.04),
          Inches(hl_top + 0.06),
          Inches(0.06),
          Inches(max(0.20, hl_h - 0.12)),
      )
      hl_l_bar.fill.solid()
      hl_l_bar.fill.fore_color.rgb = parse_color_value(hl.get("leftAccentColor"), brand_color_rgb)
      hl_l_bar.line.fill.background()
    else:
      hl_border = parse_color_value(
          hl_st.get("borderColor"),
          RGBColor(59, 130, 246) if is_dark else brand_color_rgb,
          bg_blend_rgb=blend_rgb,
      )
      hl_shape.line.color.rgb = hl_border
    tf = hl_shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18 if hl.get("hasLeftAccent") else 0.14)
    right_margin_in = 0.14
    if sub_b and sub_b.get("rect"):
      sb_w = max(0.8, sub_b["rect"].get("width", 1.2))
      right_margin_in = sb_w + 0.28
    tf.margin_right = Inches(right_margin_in)
    hl_paras = hl.get("paragraphs") or []
    if len(hl_paras) >= 2:
      for hp_idx, hp in enumerate(hl_paras):
        p = tf.paragraphs[0] if hp_idx == 0 else tf.add_paragraph()
        p.text = hp.get("text", "")
        hp_st = hp.get("styles") or hl_txt_st
        p.font.name = font_name
        p.font.size = Pt(hp_st.get("fontSizePt", 10.0))
        p.font.bold = bool(hp_st.get("isBold"))
        p.font.color.rgb = parse_color_value(hp_st.get("color"), RGBColor(17, 17, 21))
    else:
      p = tf.paragraphs[0]
      p.text = hl.get("text", "")
      p.font.name = font_name
      p.font.size = Pt(hl_txt_st.get("fontSizePt", 10.5))
      p.font.bold = True
      p.font.color.rgb = parse_color_value(hl_txt_st.get("color"), RGBColor(147, 197, 253) if is_dark else brand_color_rgb)
    if sub_b and sub_b.get("text") and sub_b.get("rect"):
      sb_r = sub_b["rect"]
      sb_st = sub_b.get("styles") or {}
      sb_w_in = max(0.75, sb_r.get("width", 1.2))
      sb_h_in = max(0.24, sb_r.get("height", 0.26))
      sb_left_in = sb_r.get("left", hl_left + hl_w - sb_w_in - 0.16)
      sb_top_in = sb_r.get("top", hl_top + (hl_h - sb_h_in) / 2.0)
      sb_shape = slide.shapes.add_shape(
          MSO_SHAPE.ROUNDED_RECTANGLE,
          Inches(sb_left_in),
          Inches(sb_top_in),
          Inches(sb_w_in),
          Inches(sb_h_in),
      )
      sb_shape.fill.solid()
      sb_shape.fill.fore_color.rgb = parse_color_value(sb_st.get("backgroundColor"), RGBColor(255, 255, 255), bg_blend_rgb=hl_bg)
      sb_border = (sb_st.get("borderColor") or "").strip()
      if sb_border and sb_border not in ("transparent", "rgba(0, 0, 0, 0)", "none"):
        sb_shape.line.color.rgb = parse_color_value(sb_border, RGBColor(191, 219, 254))
        sb_shape.line.width = Pt(1.0)
      else:
        sb_shape.line.fill.background()
      try:
        sb_shape.adjustments[0] = 0.5
      except Exception:
        pass
      sb_tf = sb_shape.text_frame
      sb_tf.word_wrap = False
      sb_tf.margin_left = Inches(0.06)
      sb_tf.margin_right = Inches(0.06)
      sb_p = sb_tf.paragraphs[0]
      sb_p.text = sub_b["text"]
      sb_p.font.name = font_name
      sb_p.font.size = Pt(sb_st.get("fontSizePt", 9.0))
      sb_p.font.bold = True
      sb_p.font.color.rgb = parse_color_value(sb_st.get("color"), brand_color_rgb)
      sb_p.alignment = PP_ALIGN.CENTER

  if geom.get("footerDividerTop"):
    f_div = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.55),
        Inches(float(geom["footerDividerTop"])),
        Inches(12.23),
        Inches(0.01),
    )
    f_div.fill.solid()
    f_div.fill.fore_color.rgb = parse_color_value(
        geom.get("footerDividerColor"),
        RGBColor(30, 41, 59) if is_dark else RGBColor(226, 232, 240),
        bg_blend_rgb=blend_rgb,
    )
    f_div.line.fill.background()

  for fn in geom.get("footnotes", []):
    fn_r = fn.get("rect", {})
    fn_st = fn.get("styles", {})
    is_right_fn = fn_r.get("left", 0.0) > 8.0 or fn_st.get("textAlign") == "right"
    fn_w_val = max(1.2, fn_r.get("width", 4.0) + 0.25)
    fn_l_val = max(0.4, (fn_r.get("left", 0.6) + fn_r.get("width", 4.0) - fn_w_val) if is_right_fn else fn_r.get("left", 0.6))
    fn_box = slide.shapes.add_textbox(
        Inches(fn_l_val),
        Inches(fn_r.get("top", 7.0)),
        Inches(fn_w_val),
        Inches(max(0.22, fn_r.get("height", 0.22))),
    )
    tf = fn_box.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    p = tf.paragraphs[0]
    p.text = fn.get("text", "")
    p.font.name = font_name
    p.font.size = Pt(fn_st.get("fontSizePt", 9.0))
    p.font.bold = bool(fn_st.get("isBold"))
    p.font.color.rgb = parse_color_value(
        fn_st.get("color"),
        RGBColor(100, 116, 139) if is_dark else RGBColor(100, 116, 139),
    )
    if is_right_fn:
      p.alignment = PP_ALIGN.RIGHT

  embed_extracted_slide_images(slide, geom, screenshot_source)


def ensure_slide_typography_consistency(
    slide,
    default_font: str = "Pretendard",
    default_color: Optional[RGBColor] = None,
    is_dark: bool = False,
) -> None:
  if default_color is None:
    default_color = RGBColor(241, 245, 249) if is_dark else RGBColor(17, 17, 21)

  for shape in slide.shapes:
    if shape.has_text_frame:
      tf = shape.text_frame
      tf.word_wrap = True
      for p in tf.paragraphs:
        if not p.font.name:
          p.font.name = default_font
        if not p.font.size:
          p.font.size = Pt(11.0)
        try:
          if not p.font.color or p.font.color.type != 1:
            p.font.color.rgb = default_color
        except Exception:
          pass
        for run in p.runs:
          if not run.font.name:
            run.font.name = default_font
          if not run.font.size:
            run.font.size = p.font.size
    elif shape.has_table:
      _disable_default_pptx_table_style(shape.table)
      for row in shape.table.rows:
        for cell in row.cells:
          cell.text_frame.word_wrap = True
          for p in cell.text_frame.paragraphs:
            if not p.font.name:
              p.font.name = default_font
            if not p.font.size:
              p.font.size = Pt(10.0)
            p_rgb = None
            try:
              if p.font.color and p.font.color.type == 1:
                p_rgb = p.font.color.rgb
              else:
                p.font.color.rgb = default_color
                p_rgb = default_color
            except Exception:
              p_rgb = default_color
            for run in p.runs:
              if not run.font.name:
                run.font.name = default_font
              if not run.font.size:
                run.font.size = p.font.size
              try:
                if not run.font.color or run.font.color.type != 1:
                  run.font.color.rgb = p_rgb
              except Exception:
                pass


def refine_card_accent_bars(slide) -> None:
  cards = [sh for sh in slide.shapes if sh.height.inches > 0.8 and sh.width.inches > 1.5 and sh.shape_type == 1]
  thin_bars = [
      sh for sh in slide.shapes
      if 0.03 <= sh.height.inches <= 0.14
      and sh.width.inches > 1.2
      and sh.top.inches < 6.5
      and not (sh.has_text_frame and sh.text_frame.text.strip())
  ]
  for b in thin_bars:
    for c in cards:
      if abs(c.left.inches - b.left.inches) < 0.45 and abs(c.top.inches - b.top.inches) < 0.14 and abs(c.width.inches - b.width.inches) < 0.8:
        inset_x = Inches(0.24)
        inset_y = Inches(0.06)
        b.left = c.left + inset_x
        b.top = c.top + inset_y
        b.width = max(Inches(0.5), c.width - (inset_x * 2))
        b.height = Inches(0.06)
        try:
          b.adjustments[0] = 0.5
        except Exception:
          pass
        try:
          b.line.fill.background()
        except Exception:
          pass

        for tb in slide.shapes:
          if tb.has_text_frame and not tb.has_table and tb != c and tb != b:
            if c.left.inches - 0.1 <= tb.left.inches <= c.left.inches + c.width.inches + 0.1:
              if c.top.inches - 0.05 <= tb.top.inches < c.top.inches + 0.16:
                tb.top = c.top + Inches(0.16)
        break


def audit_and_resolve_slide_collisions(slide) -> None:
  bg_shapes = [
      s for s in slide.shapes
      if s.shape_type == 1 and s.left.inches <= 0.1 and s.top.inches <= 0.1 and s.width.inches >= 13.0 and s.height.inches >= 7.0
  ]
  if len(bg_shapes) > 1:
    for dup in bg_shapes[1:]:
      try:
        slide.shapes._spTree.remove(dup._element)
      except Exception:
        pass

  seen_signatures = set()
  for s in list(slide.shapes):
    if s.has_text_frame and not s.has_table:
      txt = s.text_frame.text.strip()
      if txt and len(txt) > 5:
        sig = (round(s.left.inches, 2), round(s.top.inches, 2), txt)
        if sig in seen_signatures:
          try:
            slide.shapes._spTree.remove(s._element)
          except Exception:
            pass
          continue
        seen_signatures.add(sig)

  shapes = [
      s for s in slide.shapes
      if not (s.left.inches <= 0.1 and s.top.inches <= 0.1 and s.width.inches >= 13.0)
  ]

  for i in range(len(shapes)):
    s1 = shapes[i]
    l1, t1, w1, h1 = s1.left.inches, s1.top.inches, s1.width.inches, s1.height.inches
    r1, b1 = l1 + w1, t1 + h1

    for j in range(len(shapes)):
      if i == j:
        continue
      s2 = shapes[j]
      l2, t2, w2, h2 = s2.left.inches, s2.top.inches, s2.width.inches, s2.height.inches
      r2, b2 = l2 + w2, t2 + h2

      if l1 <= l2 + 0.08 and t1 <= t2 + 0.08 and r1 >= r2 - 0.30 and b1 >= b2 - 0.15:
        continue
      if l2 <= l1 + 0.08 and t2 <= t1 + 0.08 and r2 >= r1 - 0.30 and b2 >= b1 - 0.15:
        continue

      if abs(t1 - t2) > 0.12:
        continue

      x_overlap = min(r1, r2) - max(l1, l2)
      y_overlap = min(b1, b2) - max(t1, t2)

      if x_overlap > 0.02 and y_overlap > 0.05:
        if s1.shape_type == 1 and s2.shape_type == 17 and w1 < 1.8 and h1 < 0.55 and l1 < l2 and l2 > l1 + w1 * 0.4:
          new_left = r1 + 0.06
          shift = new_left - l2
          s2.left = Inches(new_left)
          s2.width = Inches(max(0.6, w2 - shift))
          l2, r2 = new_left, new_left + max(0.6, w2 - shift)

        elif s1.shape_type == 17 and s2.shape_type == 1 and w2 < 2.2 and h2 < 0.50 and l1 + 0.4 < l2:
          new_w = max(0.6, l2 - l1 - 0.06)
          s1.width = Inches(new_w)
          w1, r1 = new_w, l1 + new_w

        elif s1.shape_type == 17 and s2.shape_type == 17 and l1 + 0.4 < l2:
          new_w = max(0.6, l2 - l1 - 0.06)
          s1.width = Inches(new_w)
          w1, r1 = new_w, l1 + new_w


def harmonize_deck_presentation_fidelity(
    prs: Presentation,
    default_font: str = "Pretendard",
    default_bg_hex: Optional[str] = None,
) -> None:
  is_dark = False
  deck_bg_rgb = RGBColor(255, 255, 255)
  if default_bg_hex:
    bg_tuple = _hex_to_rgb_tuple(default_bg_hex, (255, 255, 255))
    deck_bg_rgb = RGBColor(*bg_tuple)
    is_dark = (bg_tuple[0] * 0.299 + bg_tuple[1] * 0.587 + bg_tuple[2] * 0.114) < 128
  elif len(prs.slides) > 0 and len(prs.slides[0].shapes) > 0:
    for sh in prs.slides[0].shapes:
      if sh.shape_type == 1 and sh.left.inches <= 0.1 and sh.top.inches <= 0.1 and sh.width.inches >= 13.0:
        try:
          col = sh.fill.fore_color.rgb
          deck_bg_rgb = col
          is_dark = (col[0] * 0.299 + col[1] * 0.587 + col[2] * 0.114) < 128
          break
        except Exception:
          pass

  for slide in prs.slides:
    slide_bg_rgb = deck_bg_rgb
    slide_is_dark = is_dark
    for sh in slide.shapes:
      if sh.shape_type == 1 and sh.left.inches <= 0.1 and sh.top.inches <= 0.1 and sh.width.inches >= 13.0 and sh.height.inches >= 7.0:
        try:
          col = sh.fill.fore_color.rgb
          if col is not None:
            slide_bg_rgb = col
            slide_is_dark = (col[0] * 0.299 + col[1] * 0.587 + col[2] * 0.114) < 128
            break
        except Exception:
          pass

    default_color = RGBColor(241, 245, 249) if slide_is_dark else RGBColor(17, 17, 21)
    ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=slide_is_dark)
    ensure_slide_typography_consistency(slide, default_font=default_font, default_color=default_color, is_dark=slide_is_dark)
    refine_card_accent_bars(slide)
    audit_and_resolve_slide_collisions(slide)

    for s in slide.shapes:
      if s.has_text_frame and not s.has_table:
        for p in s.text_frame.paragraphs:
          if not p.font.name:
            p.font.name = default_font
          if not p.font.size:
            p.font.size = Pt(10.5)
          for r in p.runs:
            if not r.font.name:
              r.font.name = default_font
            if not r.font.size:
              r.font.size = p.font.size


def build_native_slide_from_code(
    slide,
    prs: Presentation,
    code_str: Optional[str],
    screenshot_source: Optional[Union[bytes, str]] = None,
    design_tokens: Optional[Dict[str, Any]] = None,
    slide_dom_data: Optional[Dict[str, Any]] = None,
    slide_geometry: Optional[Dict[str, Any]] = None,
    slide_index: int = 1,
) -> None:
  font_family = (design_tokens or {}).get("font_name", "Pretendard")
  text_color_hex = (design_tokens or {}).get("text_main_hex", "#111115")
  default_text_color = _hex_to_pptx_color(text_color_hex, default=(17, 17, 21))
  brand_color = _hex_to_pptx_color((design_tokens or {}).get("brand_color_hex"), default=(37, 99, 235))

  geom_bg = (slide_geometry or {}).get("slideBgColor")
  token_bg = (design_tokens or {}).get("bg_color_hex")
  slide_bg_rgb = parse_color_value(geom_bg or token_bg, default=RGBColor(255, 255, 255))
  is_dark = (slide_bg_rgb[0] * 0.299 + slide_bg_rgb[1] * 0.587 + slide_bg_rgb[2] * 0.114) < 128
  if default_text_color == RGBColor(17, 17, 21) and is_dark:
    default_text_color = RGBColor(241, 245, 249)

  ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)

  def _build_deterministic_fallback() -> None:
    if slide_geometry:
      build_slide_from_geometry(
          slide,
          slide_geometry,
          font_name=font_family,
          brand_color_rgb=brand_color,
          screenshot_source=screenshot_source,
      )
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      refine_card_accent_bars(slide)
      audit_and_resolve_slide_collisions(slide)
      return
    if slide_dom_data and slide_dom_data.get("tables"):
      t_data = slide_dom_data["tables"][0]
      build_styled_native_table(slide, t_data, Inches(0.8), Inches(1.95), Inches(11.733), Inches(4.5), font_name=font_family)
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)

  if not code_str:
    _build_deterministic_fallback()
    return

  clean_code = code_str
  if "```python" in clean_code:
    clean_code = clean_code.split("```python", 1)[1].split("```", 1)[0].strip()
  elif "```" in clean_code:
    clean_code = clean_code.split("```", 1)[1].split("```", 1)[0].strip()

  tbl_dom = (
      slide_geometry.get("tables", [{}])[0]
      if slide_geometry and slide_geometry.get("tables")
      else (slide_dom_data.get("tables", [{}])[0] if slide_dom_data and slide_dom_data.get("tables") else None)
  )

  scope = {
      "prs": prs,
      "slide": slide,
      "slide_index": slide_index,
      "slide_geometry": slide_geometry,
      "slide_dom_data": slide_dom_data,
      "design_tokens": design_tokens,
      "table_data": tbl_dom,
      "Presentation": Presentation,
      "Inches": Inches,
      "Pt": Pt,
      "RGBColor": RGBColor,
      "MSO_SHAPE": MSO_SHAPE,
      "PP_ALIGN": PP_ALIGN,
      "MSO_ANCHOR": MSO_ANCHOR,
      "build_slide_from_geometry": build_slide_from_geometry,
      "build_styled_native_table": build_styled_native_table,
      "build_styled_native_chart": build_styled_native_chart,
      "apply_semantic_styles_to_table": apply_semantic_styles_to_table,
  }

  try:
    exec(clean_code, scope)
    fn_key = f"build_slide_{slide_index}" if f"build_slide_{slide_index}" in scope else ("build_slide" if "build_slide" in scope else None)
    if fn_key:
      func = scope[fn_key]
      sig = inspect.signature(func)
      if len(sig.parameters) == 1:
        func(slide)
      elif len(sig.parameters) == 2:
        func(prs, slide)
      else:
        func(prs, slide, slide_index)

      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)

      if tbl_dom:
        has_table_shape = any(s.has_table for s in slide.shapes)
        if not has_table_shape and tbl_dom.get("rows"):
          tbl_r = tbl_dom.get("rect", {})
          t_left = Inches(tbl_r.get("left", 0.6))
          t_top = Inches(tbl_r.get("top", 1.5))
          t_w = Inches(min(tbl_r.get("width", 12.0), 12.2))
          t_h = Inches(tbl_r.get("height", 3.0))
          tbl_copy = dict(tbl_dom)
          tbl_copy["is_dark"] = is_dark
          tbl_copy["slide_bg_rgb"] = (slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2])
          build_styled_native_table(slide, tbl_copy, t_left, t_top, t_w, t_h, font_name=font_family)
        else:
          for s in slide.shapes:
            if s.has_table:
              apply_semantic_styles_to_table(
                  s,
                  tbl_dom,
                  font_name=font_family,
                  is_dark=is_dark,
                  slide_bg_rgb=(slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2]),
              )

      if slide_geometry and slide_geometry.get("charts"):
        has_chart_shape = any(getattr(s, "has_chart", False) for s in slide.shapes)
        if not has_chart_shape:
          for ch in slide_geometry["charts"]:
            build_styled_native_chart(
                slide,
                ch,
                font_name=font_family,
                is_dark=is_dark,
                slide_bg_rgb=(slide_bg_rgb[0], slide_bg_rgb[1], slide_bg_rgb[2]),
            )

      if slide_geometry:
        embed_extracted_slide_images(slide, slide_geometry, screenshot_source)

      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      refine_card_accent_bars(slide)
      audit_and_resolve_slide_collisions(slide)
    elif len(slide.shapes) > 1:
      if slide_geometry:
        embed_extracted_slide_images(slide, slide_geometry, screenshot_source)
      ensure_slide_canvas_background(slide, prs=prs, bg_color=slide_bg_rgb, is_dark=is_dark)
      ensure_slide_typography_consistency(slide, default_font=font_family, default_color=default_text_color, is_dark=is_dark)
      refine_card_accent_bars(slide)
      audit_and_resolve_slide_collisions(slide)
    else:
      _build_deterministic_fallback()
  except Exception as e:
    logger.warning("[Execution Error in custom slide builder code]: %s", e)
    _build_deterministic_fallback()
